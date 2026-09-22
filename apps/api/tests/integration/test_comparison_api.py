import io
import json
from datetime import UTC, datetime, timedelta
from uuid import UUID

import psycopg
import pytest
from argon2 import PasswordHasher
from cryptography.fernet import Fernet
from docx import Document
from httpx import ASGITransport, AsyncClient

from adapters.object_storage.client import ObjectStorageClient
from apps.worker.src.patent_evidence_worker.parse_worker import ParseTaskWorker
from patent_evidence_api.core.database import (
    create_engine,
    create_worker_session_factory,
)
from patent_evidence_api.main import create_app
from test_support import api_settings, login_with_totp, postgres_url

ORG_A = UUID("90000000-0000-4000-8000-000000000001")
ORG_B = UUID("90000000-0000-4000-8000-000000000002")
ADMIN_A = UUID("91000000-0000-4000-8000-000000000001")
ADMIN_B = UUID("91000000-0000-4000-8000-000000000002")
MEMBER_A = UUID("92000000-0000-4000-8000-000000000001")
MEMBER_B = UUID("92000000-0000-4000-8000-000000000002")
PASSWORD = "Correct horse battery staple 42"
MFA_KEY = Fernet.generate_key().decode()


class Clock:
    def __init__(self) -> None:
        self.value = datetime.now(UTC).replace(microsecond=0)

    def __call__(self) -> datetime:
        return self.value


def _seed_org_and_identities(migration_url: str, now: datetime) -> None:
    password_hash = PasswordHasher().hash(PASSWORD)
    with psycopg.connect(migration_url, autocommit=True) as conn:
        with conn.cursor() as cur:
            cur.execute(
                """INSERT INTO global_identities
                (id, email_normalized, status, security_version, created_at, updated_at)
                VALUES
                (%s, 'admin_a_comp@tenant.com', 'active', 1, %s, %s),
                (%s, 'admin_b_comp@tenant.com', 'active', 1, %s, %s)
                ON CONFLICT (id) DO NOTHING""",
                (ADMIN_A, now, now, ADMIN_B, now, now),
            )
            cur.execute(
                """INSERT INTO credentials
                (id, identity_id, credential_type, secret_hash, status, created_at, updated_at)
                VALUES
                (gen_random_uuid(), %s, 'password', %s, 'active', %s, %s),
                (gen_random_uuid(), %s, 'password', %s, 'active', %s, %s)
                ON CONFLICT DO NOTHING""",
                (ADMIN_A, password_hash, now, now, ADMIN_B, password_hash, now, now),
            )
            cur.execute(
                """INSERT INTO organizations
                (id, slug, display_name, persisted_status, created_at, updated_at)
                VALUES
                (%s, 'comp-org-a', '比对机构A', 'active', %s, %s),
                (%s, 'comp-org-b', '比对机构B', 'active', %s, %s)
                ON CONFLICT (id) DO NOTHING""",
                (ORG_A, now, now, ORG_B, now, now),
            )
            cur.execute(
                """INSERT INTO organization_plan_quotas
                (id, organization_id, plan_key, monthly_case_allowance, current_period_start, current_period_end, quota_status, created_at, updated_at)
                VALUES
                (gen_random_uuid(), %s, 'standard', 50, %s, %s, 'active', %s, %s),
                (gen_random_uuid(), %s, 'standard', 50, %s, %s, 'active', %s, %s)
                ON CONFLICT DO NOTHING""",
                (
                    ORG_A,
                    now - timedelta(days=1),
                    now + timedelta(days=29),
                    now,
                    now,
                    ORG_B,
                    now - timedelta(days=1),
                    now + timedelta(days=29),
                    now,
                    now,
                ),
            )
            cur.execute(
                """INSERT INTO organization_memberships
                (id, organization_id, identity_id, role, status, created_at, updated_at)
                VALUES
                (%s, %s, %s, 'organization_admin', 'active', %s, %s),
                (%s, %s, %s, 'organization_admin', 'active', %s, %s)
                ON CONFLICT (id) DO NOTHING""",
                (MEMBER_A, ORG_A, ADMIN_A, now, now, MEMBER_B, ORG_B, ADMIN_B, now, now),
            )


@pytest.mark.asyncio
async def test_claim_feature_comparison_matrix_lifecycle() -> None:
    migration_url = postgres_url("PE_MIGRATION_DATABASE_URL", "patent_evidence_migration")
    clock = Clock()
    _seed_org_and_identities(migration_url, clock())

    settings = api_settings(mfa_encryption_key=MFA_KEY)
    app = create_app(settings=settings, clock=clock)

    async with AsyncClient(
        transport=ASGITransport(app=app),
        base_url="http://testserver",
    ) as client:
        # 1. Login Admin A
        await login_with_totp(client, "admin_a_comp@tenant.com", PASSWORD)

        # 2. Create Case & Document & Parse
        res_case = await client.post(
            f"/api/v1/organizations/{ORG_A}/cases",
            json={
                "case_number": "2026-COMP-001",
                "title": "大模型稀疏矩阵量化推理方法",
                "technical_field": "人工智能",
            },
        )
        assert res_case.status_code == 201
        case_id = res_case.json()["case"]["id"]

        doc = Document()
        doc.add_heading("技术领域", level=1)
        doc.add_paragraph("本申请涉及一种面向大模型混合精度量化的推理加速方法。")
        doc.add_heading("发明内容", level=1)
        doc.add_paragraph("本发明通过获取权重张量并执行奇异值分解与稀疏量化查找表构建。")
        buf = io.BytesIO()
        doc.save(buf)

        import base64
        await client.post(
            f"/api/v1/organizations/{ORG_A}/cases/{case_id}/documents",
            json={
                "filename": "disclosure.docx",
                "content_base64": base64.b64encode(buf.getvalue()).decode("ascii"),
                "content_type": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            },
        )

        worker_engine = create_engine(settings.worker_database_url)
        worker_session_factory = create_worker_session_factory(worker_engine)
        storage = ObjectStorageClient()
        worker = ParseTaskWorker(worker_session_factory, storage, clock=clock)
        await worker.run_once()
        await worker_engine.dispose()

        # Confirm Document Version
        case_detail = (await client.get(f"/api/v1/organizations/{ORG_A}/cases/{case_id}")).json()
        doc_ver_id = case_detail["document_version"]["id"]
        await client.post(
            f"/api/v1/organizations/{ORG_A}/cases/{case_id}/document-versions/{doc_ver_id}/confirm"
        )

        # Extract & Confirm Features
        res_feats = await client.post(
            f"/api/v1/organizations/{ORG_A}/cases/{case_id}/features/extract"
        )
        f_ver_id = res_feats.json()["version"]["id"]
        await client.post(
            f"/api/v1/organizations/{ORG_A}/cases/{case_id}/features/versions/{f_ver_id}/confirm"
        )

        # Generate Strategy & Import Candidate D1
        await client.post(
            f"/api/v1/organizations/{ORG_A}/cases/{case_id}/search/strategies/generate"
        )
        cnipr_csv = """公开号,发明名称,摘要,公开日,申请人,IPC分类号
CN117283912A,大模型量化加速系统,通过奇异值分解优化权重分布降低张量冗余,2024-03-15,前沿智能,G06N 3/08
"""
        await client.post(
            f"/api/v1/organizations/{ORG_A}/cases/{case_id}/search/candidates/import-cnipr",
            json={"raw_content": cnipr_csv},
        )

        # Mark candidate as included
        cands = (await client.get(f"/api/v1/organizations/{ORG_A}/cases/{case_id}/search/candidates")).json()["items"]
        cand_id = cands[0]["id"]
        await client.post(
            f"/api/v1/organizations/{ORG_A}/cases/{case_id}/search/candidates/{cand_id}/triage",
            json={"triage_status": "included"},
        )

        # 3. Generate Feature Comparison Matrix
        res_gen_matrix = await client.post(
            f"/api/v1/organizations/{ORG_A}/cases/{case_id}/comparisons/generate"
        )
        assert res_gen_matrix.status_code == 201
        matrix_data = res_gen_matrix.json()
        matrix_id = matrix_data["matrix"]["id"]
        assert matrix_data["matrix"]["status"] == "draft"
        assert len(matrix_data["features"]) >= 2
        assert len(matrix_data["candidates"]) >= 1
        assert len(matrix_data["comparisons"]) >= 2

        comp_item = matrix_data["comparisons"][0]
        assert comp_item["judgment"] in ("identical", "equivalent", "different", "insufficient_evidence")
        assert comp_item["evidence_status"] in ("abstract_only", "missing_source_text")
        assert comp_item["evaluation_source"] == "deterministic_baseline"

        # 4. Update Comparison Item (Manual Override)
        res_update_comp = await client.put(
            f"/api/v1/organizations/{ORG_A}/cases/{case_id}/comparisons/items/{comp_item['id']}",
            json={
                "judgment": "equivalent",
                "citation_location": "说明书第[0032]段",
                "citation_quote": "该系统在解码阶段采用低位宽映射实现加速",
                "reasoning_analysis": "人工审查意见：对比文件公开了等同的稀疏查找表映射技术手段。",
            },
        )
        assert res_update_comp.status_code == 200
        updated_comp = [c for c in res_update_comp.json()["comparisons"] if c["id"] == comp_item["id"]][0]
        assert updated_comp["judgment"] == "equivalent"
        assert updated_comp["is_manually_edited"] is True
        assert "人工审查意见" in updated_comp["reasoning_analysis"]

        # 5. Confirm and Lock Comparison Matrix
        res_confirm_matrix = await client.post(
            f"/api/v1/organizations/{ORG_A}/cases/{case_id}/comparisons/{matrix_id}/confirm"
        )
        assert res_confirm_matrix.status_code == 200
        assert res_confirm_matrix.json()["matrix"]["status"] == "confirmed"

        # Check Case Status advanced to comparisons_confirmed
        case_after = (await client.get(f"/api/v1/organizations/{ORG_A}/cases/{case_id}")).json()
        assert case_after["status"] == "comparisons_confirmed"

        # 6. Attempting to edit locked comparison matrix is rejected
        res_locked_edit = await client.put(
            f"/api/v1/organizations/{ORG_A}/cases/{case_id}/comparisons/items/{comp_item['id']}",
            json={"judgment": "different"},
        )
        assert res_locked_edit.status_code == 400

        # 7. Cross-Tenant Isolation: Admin B cannot access Org A's comparison matrix
        await login_with_totp(client, "admin_b_comp@tenant.com", PASSWORD)
        res_cross = await client.get(
            f"/api/v1/organizations/{ORG_B}/cases/{case_id}/comparisons/active"
        )
        assert res_cross.status_code == 404
