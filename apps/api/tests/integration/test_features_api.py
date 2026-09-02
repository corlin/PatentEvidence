import io
import json
from datetime import UTC, datetime, timedelta
from uuid import UUID

import psycopg
import pytest
from argon2 import PasswordHasher
from cryptography.fernet import Fernet
from docx import Document
from httpx import ASGITransport, AsyncClient

from adapters.object_storage.client import ObjectStorageClient
from apps.worker.src.patent_evidence_worker.parse_worker import ParseTaskWorker
from patent_evidence_api.core.database import (
    create_engine,
    create_worker_session_factory,
)
from patent_evidence_api.main import create_app
from test_support import api_settings, login_with_totp, postgres_url

ORG_A = UUID("70000000-0000-4000-8000-000000000001")
ORG_B = UUID("70000000-0000-4000-8000-000000000002")
ADMIN_A = UUID("71000000-0000-4000-8000-000000000001")
ADMIN_B = UUID("71000000-0000-4000-8000-000000000002")
MEMBER_A = UUID("72000000-0000-4000-8000-000000000001")
MEMBER_B = UUID("72000000-0000-4000-8000-000000000002")
PASSWORD = "Correct horse battery staple 42"
MFA_KEY = Fernet.generate_key().decode()


class Clock:
    def __init__(self) -> None:
        self.value = datetime.now(UTC).replace(microsecond=0)

    def __call__(self) -> datetime:
        return self.value


def _seed_org_and_identities(migration_url: str, now: datetime) -> None:
    password_hash = PasswordHasher().hash(PASSWORD)
    with psycopg.connect(migration_url, autocommit=True) as conn:
        with conn.cursor() as cur:
            cur.execute(
                """INSERT INTO global_identities
                (id, email_normalized, status, security_version, created_at, updated_at)
                VALUES
                (%s, 'admin_a_feat@tenant.com', 'active', 1, %s, %s),
                (%s, 'admin_b_feat@tenant.com', 'active', 1, %s, %s)
                ON CONFLICT (id) DO NOTHING""",
                (ADMIN_A, now, now, ADMIN_B, now, now),
            )
            cur.execute(
                """INSERT INTO credentials
                (id, identity_id, credential_type, secret_hash, status, created_at, updated_at)
                VALUES
                (gen_random_uuid(), %s, 'password', %s, 'active', %s, %s),
                (gen_random_uuid(), %s, 'password', %s, 'active', %s, %s)
                ON CONFLICT DO NOTHING""",
                (ADMIN_A, password_hash, now, now, ADMIN_B, password_hash, now, now),
            )
            cur.execute(
                """INSERT INTO organizations
                (id, slug, display_name, persisted_status, created_at, updated_at)
                VALUES
                (%s, 'feat-org-a', '特征机构A', 'active', %s, %s),
                (%s, 'feat-org-b', '特征机构B', 'active', %s, %s)
                ON CONFLICT (id) DO NOTHING""",
                (ORG_A, now, now, ORG_B, now, now),
            )
            cur.execute(
                """INSERT INTO organization_plan_quotas
                (id, organization_id, plan_key, monthly_case_allowance, current_period_start, current_period_end, quota_status, created_at, updated_at)
                VALUES
                (gen_random_uuid(), %s, 'standard', 50, %s, %s, 'active', %s, %s),
                (gen_random_uuid(), %s, 'standard', 50, %s, %s, 'active', %s, %s)
                ON CONFLICT DO NOTHING""",
                (
                    ORG_A,
                    now - timedelta(days=1),
                    now + timedelta(days=29),
                    now,
                    now,
                    ORG_B,
                    now - timedelta(days=1),
                    now + timedelta(days=29),
                    now,
                    now,
                ),
            )
            cur.execute(
                """INSERT INTO organization_memberships
                (id, organization_id, identity_id, role, status, created_at, updated_at)
                VALUES
                (%s, %s, %s, 'organization_admin', 'active', %s, %s),
                (%s, %s, %s, 'organization_admin', 'active', %s, %s)
                ON CONFLICT (id) DO NOTHING""",
                (MEMBER_A, ORG_A, ADMIN_A, now, now, MEMBER_B, ORG_B, ADMIN_B, now, now),
            )


@pytest.mark.asyncio
async def test_feature_modeling_lifecycle_and_confirmation() -> None:
    migration_url = postgres_url("PE_MIGRATION_DATABASE_URL", "patent_evidence_migration")
    clock = Clock()
    _seed_org_and_identities(migration_url, clock())

    settings = api_settings(mfa_encryption_key=MFA_KEY)
    app = create_app(settings=settings, clock=clock)

    async with AsyncClient(
        transport=ASGITransport(app=app),
        base_url="http://testserver",
    ) as client:
        # 1. Login Admin A
        await login_with_totp(client, "admin_a_feat@tenant.com", PASSWORD)

        # 2. Create Case & Document
        res_case = await client.post(
            f"/api/v1/organizations/{ORG_A}/cases",
            json={
                "case_number": "2026-FEAT-001",
                "title": "大模型混合量化技术",
                "technical_field": "人工智能",
            },
        )
        assert res_case.status_code == 201
        case_id = res_case.json()["case"]["id"]

        # 3. Create document directly and parse
        doc = Document()
        doc.add_heading("技术领域", level=1)
        doc.add_paragraph("本申请涉及一种面向边缘设备的大模型多尺度量化方法。")
        doc.add_heading("发明内容", level=1)
        doc.add_paragraph("本发明通过获取权重张量并执行对角加权分解；基于通道活跃度构建稀疏量化查找表。")
        buf = io.BytesIO()
        doc.save(buf)
        docx_bytes = buf.getvalue()

        import base64
        await client.post(
            f"/api/v1/organizations/{ORG_A}/cases/{case_id}/documents",
            json={
                "filename": "disclosure.docx",
                "content_base64": base64.b64encode(docx_bytes).decode("ascii"),
                "content_type": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            },
        )

        worker_engine = create_engine(settings.worker_database_url)
        worker_session_factory = create_worker_session_factory(worker_engine)
        storage = ObjectStorageClient()
        worker = ParseTaskWorker(worker_session_factory, storage, clock=clock)
        await worker.run_once()
        await worker_engine.dispose()

        # Confirm document version
        case_detail = (await client.get(f"/api/v1/organizations/{ORG_A}/cases/{case_id}")).json()
        doc_ver_id = case_detail["document_version"]["id"]
        await client.post(
            f"/api/v1/organizations/{ORG_A}/cases/{case_id}/document-versions/{doc_ver_id}/confirm"
        )

        # 4. Trigger feature draft extraction
        res_extract = await client.post(
            f"/api/v1/organizations/{ORG_A}/cases/{case_id}/features/extract"
        )
        assert res_extract.status_code == 201
        feat_data = res_extract.json()
        version_id = feat_data["version"]["id"]
        assert feat_data["version"]["status"] == "draft"
        assert feat_data["version"]["version_number"] == 1
        assert len(feat_data["features"]) >= 2
        f1 = feat_data["features"][0]
        f2 = feat_data["features"][1]
        assert f1["source_paragraph_id"] is not None

        # 5. Update Feature 1
        res_up = await client.put(
            f"/api/v1/organizations/{ORG_A}/cases/{case_id}/features/versions/{version_id}/items/{f1['id']}",
            json={
                "feature_statement": "一种面向边缘设备的大模型多尺度量化方法，其特征在于：",
                "feature_type": "preamble",
            },
        )
        assert res_up.status_code == 200
        assert res_up.json()["features"][0]["feature_statement"] == "一种面向边缘设备的大模型多尺度量化方法，其特征在于："

        # 6. Split Feature 2 (1 into 2)
        res_split = await client.post(
            f"/api/v1/organizations/{ORG_A}/cases/{case_id}/features/versions/{version_id}/items/{f2['id']}/split",
            json={
                "part1_statement": "步骤一：获取权重张量并执行对角加权分解",
                "part2_statement": "步骤二：基于通道活跃度构建稀疏量化查找表",
            },
        )
        assert res_split.status_code == 200
        features_after_split = res_split.json()["features"]
        assert len(features_after_split) >= 3

        # 7. Add a new custom feature
        res_add = await client.post(
            f"/api/v1/organizations/{ORG_A}/cases/{case_id}/features/versions/{version_id}/items",
            json={
                "feature_code": "F4",
                "feature_type": "dependent",
                "feature_statement": "根据权利要求 1 所述的方法，其中加权系数为动态自适应设定。",
                "source_paragraph_id": "p2",
                "citation_quote": "动态自适应设定",
            },
        )
        assert res_add.status_code == 201

        # 8. Merge two features
        current_features = res_add.json()["features"]
        target1 = current_features[1]["id"]
        target2 = current_features[2]["id"]
        res_merge = await client.post(
            f"/api/v1/organizations/{ORG_A}/cases/{case_id}/features/versions/{version_id}/merge",
            json={
                "feature_id_1": target1,
                "feature_id_2": target2,
                "merged_statement": "合并步骤：获取权重张量并构建稀疏量化查找表",
            },
        )
        assert res_merge.status_code == 200

        # 9. Confirm and Lock Feature Set Version
        res_confirm = await client.post(
            f"/api/v1/organizations/{ORG_A}/cases/{case_id}/features/versions/{version_id}/confirm"
        )
        assert res_confirm.status_code == 200
        assert res_confirm.json()["version"]["status"] == "confirmed"
        assert res_confirm.json()["version"]["confirmed_at"] is not None

        # Check Case Status advanced to features_confirmed
        case_res_after = await client.get(f"/api/v1/organizations/{ORG_A}/cases/{case_id}")
        assert case_res_after.json()["status"] == "features_confirmed"

        # 10. Attempting to modify confirmed version is rejected (immutable)
        res_forbidden_edit = await client.put(
            f"/api/v1/organizations/{ORG_A}/cases/{case_id}/features/versions/{version_id}/items/{f1['id']}",
            json={"feature_statement": "恶意修改已确认版本"},
        )
        assert res_forbidden_edit.status_code == 400

        # 11. Create Revision (clones v1 to draft v2)
        res_rev = await client.post(
            f"/api/v1/organizations/{ORG_A}/cases/{case_id}/features/versions/{version_id}/revision"
        )
        assert res_rev.status_code == 200
        v2_data = res_rev.json()
        assert v2_data["version"]["version_number"] == 2
        assert v2_data["version"]["status"] == "draft"
        assert v2_data["version"]["parent_version_id"] == version_id
        assert len(v2_data["features"]) > 0

        # 12. Cross-tenant access: Admin B cannot view or modify Org A's feature sets
        await login_with_totp(client, "admin_b_feat@tenant.com", PASSWORD)
        res_cross = await client.get(
            f"/api/v1/organizations/{ORG_B}/cases/{case_id}/features/active"
        )
        assert res_cross.status_code == 404
