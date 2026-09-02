import base64
import io
import json
from datetime import UTC, datetime, timedelta
from uuid import UUID, uuid4

import psycopg
import pytest
from argon2 import PasswordHasher
from cryptography.fernet import Fernet
from docx import Document
from httpx import ASGITransport, AsyncClient

from adapters.object_storage.client import ObjectStorageClient
from apps.worker.src.patent_evidence_worker.parse_worker import ParseTaskWorker
from patent_evidence_api.core.database import (
    create_engine,
    create_worker_session_factory,
)
from patent_evidence_api.main import create_app
from test_support import api_settings, login_with_totp, postgres_url

SUPERADMIN_ID = UUID("80000000-0000-4000-8000-000000000001")
PASSWORD = "Correct horse battery staple 42"
MFA_KEY = Fernet.generate_key().decode()


class Clock:
    def __init__(self) -> None:
        self.value = datetime.now(UTC).replace(microsecond=0)

    def __call__(self) -> datetime:
        return self.value


def _seed_superadmin(migration_url: str, now: datetime) -> None:
    password_hash = PasswordHasher().hash(PASSWORD)
    with psycopg.connect(migration_url, autocommit=True) as conn:
        with conn.cursor() as cur:
            cur.execute(
                """INSERT INTO global_identities
                (id, email_normalized, status, security_version, created_at, updated_at)
                VALUES
                (%s, 'superadmin_e2e@platform.com', 'active', 1, %s, %s)
                ON CONFLICT (id) DO NOTHING""",
                (SUPERADMIN_ID, now, now),
            )
            cur.execute(
                """INSERT INTO credentials
                (id, identity_id, credential_type, secret_hash, status, created_at, updated_at)
                VALUES
                (gen_random_uuid(), %s, 'password', %s, 'active', %s, %s)
                ON CONFLICT DO NOTHING""",
                (SUPERADMIN_ID, password_hash, now, now),
            )


@pytest.mark.asyncio
async def test_p0_full_lifecycle_e2e_integration() -> None:
    migration_url = postgres_url("PE_MIGRATION_DATABASE_URL", "patent_evidence_migration")
    clock = Clock()
    _seed_superadmin(migration_url, clock())

    settings = api_settings(mfa_encryption_key=MFA_KEY)
    app = create_app(settings=settings, clock=clock)

    async with AsyncClient(
        transport=ASGITransport(app=app),
        base_url="http://testserver",
    ) as client:
        # Step 1: Login Superadmin
        await login_with_totp(client, "superadmin_e2e@platform.com", PASSWORD)

        # Step 2: Superadmin provisions Organization A
        res_prov = await client.post(
            "/api/v1/platform/organizations",
            json={
                "slug": "final-e2e-org-a",
                "display_name": "终审验收代理机构A",
                "admin_email": "tenant_admin_a@firm.com",
                "monthly_case_allowance": 100,
            },
        )
        assert res_prov.status_code == 201
        prov_data = res_prov.json()
        org_a_id = prov_data["organization"]["id"]
        invitation_token = prov_data["invitation"]["token"]

        # Step 3: Tenant Admin A accepts invitation and sets password
        res_inspect = await client.post(
            "/api/v1/invitations/inspect",
            json={"token": invitation_token},
        )
        assert res_inspect.status_code == 200

        res_accept = await client.post(
            "/api/v1/invitations/accept",
            json={
                "token": invitation_token,
                "display_name": "张合伙人",
                "password": PASSWORD,
            },
        )
        assert res_accept.status_code == 200

        # Step 4: Login as Tenant Admin A with TOTP
        await login_with_totp(client, "tenant_admin_a@firm.com", PASSWORD)

        # Step 5: Create Patent Case
        res_case = await client.post(
            f"/api/v1/organizations/{org_a_id}/cases",
            json={
                "case_number": "2026-FINAL-E2E-001",
                "title": "一种面向低位宽大语言模型的稀疏量化加速系统与方法",
                "technical_field": "人工智能与大模型加速",
            },
        )
        assert res_case.status_code == 201
        case_id = res_case.json()["case"]["id"]

        # Step 6: Upload Technical Disclosure Document & Run Parse Worker
        doc = Document()
        doc.add_heading("技术领域", level=1)
        doc.add_paragraph("本申请涉及一种面向大语言模型混合精度量化的推理加速方法。")
        doc.add_heading("发明内容", level=1)
        doc.add_paragraph("本发明通过获取权重张量并执行奇异值分解与稀疏量化查找表构建，从而降低量化误差并提高GPU前向计算吞吐。")
        buf = io.BytesIO()
        doc.save(buf)

        res_upload = await client.post(
            f"/api/v1/organizations/{org_a_id}/cases/{case_id}/documents",
            json={
                "filename": "disclosure_full_e2e.docx",
                "content_base64": base64.b64encode(buf.getvalue()).decode("ascii"),
                "content_type": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            },
        )
        assert res_upload.status_code == 201

        # Run Worker
        worker_engine = create_engine(settings.worker_database_url)
        worker_session_factory = create_worker_session_factory(worker_engine)
        storage = ObjectStorageClient()
        worker = ParseTaskWorker(worker_session_factory, storage, clock=clock)
        await worker.run_once()
        await worker_engine.dispose()

        # Step 7: Confirm Document Version
        case_detail = (await client.get(f"/api/v1/organizations/{org_a_id}/cases/{case_id}")).json()
        doc_ver_id = case_detail["document_version"]["id"]
        res_conf_doc = await client.post(
            f"/api/v1/organizations/{org_a_id}/cases/{case_id}/document-versions/{doc_ver_id}/confirm"
        )
        assert res_conf_doc.status_code == 200

        # Step 8: Extract and Confirm Features
        res_ext = await client.post(
            f"/api/v1/organizations/{org_a_id}/cases/{case_id}/features/extract"
        )
        assert res_ext.status_code == 201
        fsv_id = res_ext.json()["version"]["id"]
        res_conf_feat = await client.post(
            f"/api/v1/organizations/{org_a_id}/cases/{case_id}/features/versions/{fsv_id}/confirm"
        )
        assert res_conf_feat.status_code == 200

        # Step 9: Search Strategy Planning & CNIPR Handoff Package
        res_strat = await client.post(
            f"/api/v1/organizations/{org_a_id}/cases/{case_id}/search/strategies/generate"
        )
        assert res_strat.status_code == 201
        res_handoff = await client.get(
            f"/api/v1/organizations/{org_a_id}/cases/{case_id}/search/handoff"
        )
        assert res_handoff.status_code == 200
        assert "CNIPR 官方专利检索人工交接规范包" in res_handoff.json()["markdown"]

        # Step 10: Import CNIPR Candidate Patents
        cnipr_csv = """公开号,发明名称,摘要,公开日,申请人,IPC分类号
CN119999999A,大模型稀疏量化系统,基于奇异值分解优化矩阵权重分布降低计算冗余,2024-05-10,前沿科技研究院,G06N 3/08
CN118888888A,机械传动系统,一种高耐磨圆柱齿轮与箱体结构,2024-02-01,工业制造厂,F16H 1/00
"""
        res_import = await client.post(
            f"/api/v1/organizations/{org_a_id}/cases/{case_id}/search/candidates/import-cnipr",
            json={"raw_content": cnipr_csv},
        )
        assert res_import.status_code == 201

        # Step 11: Candidate Pool Triage
        cands = (await client.get(f"/api/v1/organizations/{org_a_id}/cases/{case_id}/search/candidates")).json()["items"]
        assert len(cands) >= 2
        d1 = [c for c in cands if c["publication_number"] == "CN119999999A"][0]
        d2 = [c for c in cands if c["publication_number"] == "CN118888888A"][0]

        # Include D1, Exclude D2 with reason
        await client.post(
            f"/api/v1/organizations/{org_a_id}/cases/{case_id}/search/candidates/{d1['id']}/triage",
            json={"triage_status": "included"},
        )
        await client.post(
            f"/api/v1/organizations/{org_a_id}/cases/{case_id}/search/candidates/{d2['id']}/triage",
            json={"triage_status": "excluded", "exclusion_reason": "技术领域不相关 (机械结构 vs 算法加速)"},
        )

        # Step 12: Dual-Mode Comparison Engine Execution & Manual Cell Edit
        res_comp_gen = await client.post(
            f"/api/v1/organizations/{org_a_id}/cases/{case_id}/comparisons/generate"
        )
        assert res_comp_gen.status_code == 201
        matrix_data = res_comp_gen.json()
        matrix_id = matrix_data["matrix"]["id"]
        comp_item = matrix_data["comparisons"][0]

        # Manual edit cell
        await client.put(
            f"/api/v1/organizations/{org_a_id}/cases/{case_id}/comparisons/items/{comp_item['id']}",
            json={
                "judgment": "equivalent",
                "citation_location": "说明书第[0028]段",
                "citation_quote": "采用奇异值分解优化矩阵权重分布降低计算冗余",
                "reasoning_analysis": "人工终审意见：对比文件公开了等同技术手段，二者解决的技术问题与效果实质相同。",
            },
        )

        # Step 13: Confirm & Lock Comparison Matrix
        res_conf_mat = await client.post(
            f"/api/v1/organizations/{org_a_id}/cases/{case_id}/comparisons/{matrix_id}/confirm"
        )
        assert res_conf_mat.status_code == 200

        # Step 14: Seal Immutable Evidence Snapshot & Compute Root SHA-256
        res_seal = await client.post(
            f"/api/v1/organizations/{org_a_id}/cases/{case_id}/evidence/seal"
        )
        assert res_seal.status_code == 201
        seal_data = res_seal.json()
        root_sha256 = seal_data["snapshot"]["root_sha256"]
        assert len(root_sha256) == 64
        assert seal_data["snapshot"]["status"] == "sealed"

        # Step 15: Verify Report Generated
        report_content = seal_data["report"]["content"]
        assert "# 专利证据分析与法律评估报告" in report_content
        assert "2026-FINAL-E2E-001" in report_content
        assert root_sha256 in report_content

        # Step 16: Verify Case Status is Completed
        final_case = (await client.get(f"/api/v1/organizations/{org_a_id}/cases/{case_id}")).json()
        assert final_case["status"] == "completed"

        # Step 17: Cross-Tenant Isolation Security Assertion
        # Provision Organization B and attempt to access Org A's case
        await login_with_totp(client, "superadmin_e2e@platform.com", PASSWORD)
        res_prov_b = await client.post(
            "/api/v1/platform/organizations",
            json={
                "slug": "final-e2e-org-b",
                "display_name": "机构B",
                "admin_email": "tenant_admin_b@firm.com",
                "monthly_case_allowance": 100,
            },
        )
        inv_token_b = res_prov_b.json()["invitation"]["token"]
        await client.post(
            "/api/v1/invitations/accept",
            json={"token": inv_token_b, "display_name": "李合伙人", "password": PASSWORD},
        )
        await login_with_totp(client, "tenant_admin_b@firm.com", PASSWORD)
        org_b_id = res_prov_b.json()["organization"]["id"]

        # Org B attempting to read Org A's case or reports MUST return 404
        assert (await client.get(f"/api/v1/organizations/{org_b_id}/cases/{case_id}")).status_code == 404
        assert (await client.get(f"/api/v1/organizations/{org_b_id}/cases/{case_id}/evidence/active")).status_code == 404
        assert (await client.get(f"/api/v1/organizations/{org_b_id}/cases/{case_id}/reports/active")).status_code == 404
