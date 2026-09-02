import io
import json
from datetime import UTC, datetime, timedelta
from uuid import UUID

import psycopg
import pytest
from argon2 import PasswordHasher
from cryptography.fernet import Fernet
from docx import Document
from httpx import ASGITransport, AsyncClient

from adapters.object_storage.client import ObjectStorageClient
from apps.worker.src.patent_evidence_worker.parse_worker import ParseTaskWorker
from patent_evidence_api.core.database import (
    create_engine,
    create_worker_session_factory,
)
from patent_evidence_api.main import create_app
from test_support import api_settings, login_with_totp, postgres_url

ORG_A = UUID("90000000-0000-4000-8000-000000000001")
ORG_B = UUID("90000000-0000-4000-8000-000000000002")
ADMIN_A = UUID("91000000-0000-4000-8000-000000000001")
ADMIN_B = UUID("91000000-0000-4000-8000-000000000002")
MEMBER_A = UUID("92000000-0000-4000-8000-000000000001")
MEMBER_B = UUID("92000000-0000-4000-8000-000000000002")
PASSWORD = "Correct horse battery staple 42"
MFA_KEY = Fernet.generate_key().decode()


class Clock:
    def __init__(self) -> None:
        self.value = datetime.now(UTC).replace(microsecond=0)

    def __call__(self) -> datetime:
        return self.value


def _seed_org_and_identities(migration_url: str, now: datetime) -> None:
    password_hash = PasswordHasher().hash(PASSWORD)
    with psycopg.connect(migration_url, autocommit=True) as conn:
        with conn.cursor() as cur:
            cur.execute(
                """INSERT INTO global_identities
                (id, email_normalized, status, security_version, created_at, updated_at)
                VALUES
                (%s, 'admin_a_rep@tenant.com', 'active', 1, %s, %s),
                (%s, 'admin_b_rep@tenant.com', 'active', 1, %s, %s)
                ON CONFLICT (id) DO NOTHING""",
                (ADMIN_A, now, now, ADMIN_B, now, now),
            )
            cur.execute(
                """INSERT INTO credentials
                (id, identity_id, credential_type, secret_hash, status, created_at, updated_at)
                VALUES
                (gen_random_uuid(), %s, 'password', %s, 'active', %s, %s),
                (gen_random_uuid(), %s, 'password', %s, 'active', %s, %s)
                ON CONFLICT DO NOTHING""",
                (ADMIN_A, password_hash, now, now, ADMIN_B, password_hash, now, now),
            )
            cur.execute(
                """INSERT INTO organizations
                (id, slug, display_name, persisted_status, created_at, updated_at)
                VALUES
                (%s, 'rep-org-a', '报告机构A', 'active', %s, %s),
                (%s, 'rep-org-b', '报告机构B', 'active', %s, %s)
                ON CONFLICT (id) DO NOTHING""",
                (ORG_A, now, now, ORG_B, now, now),
            )
            cur.execute(
                """INSERT INTO organization_plan_quotas
                (id, organization_id, plan_key, monthly_case_allowance, current_period_start, current_period_end, quota_status, created_at, updated_at)
                VALUES
                (gen_random_uuid(), %s, 'standard', 50, %s, %s, 'active', %s, %s),
                (gen_random_uuid(), %s, 'standard', 50, %s, %s, 'active', %s, %s)
                ON CONFLICT DO NOTHING""",
                (
                    ORG_A,
                    now - timedelta(days=1),
                    now + timedelta(days=29),
                    now,
                    now,
                    ORG_B,
                    now - timedelta(days=1),
                    now + timedelta(days=29),
                    now,
                    now,
                ),
            )
            cur.execute(
                """INSERT INTO organization_memberships
                (id, organization_id, identity_id, role, status, created_at, updated_at)
                VALUES
                (%s, %s, %s, 'organization_admin', 'active', %s, %s),
                (%s, %s, %s, 'organization_admin', 'active', %s, %s)
                ON CONFLICT (id) DO NOTHING""",
                (MEMBER_A, ORG_A, ADMIN_A, now, now, MEMBER_B, ORG_B, ADMIN_B, now, now),
            )


@pytest.mark.asyncio
async def test_evidence_sealing_and_report_generation_lifecycle() -> None:
    migration_url = postgres_url("PE_MIGRATION_DATABASE_URL", "patent_evidence_migration")
    clock = Clock()
    _seed_org_and_identities(migration_url, clock())

    settings = api_settings(mfa_encryption_key=MFA_KEY)
    app = create_app(settings=settings, clock=clock)

    async with AsyncClient(
        transport=ASGITransport(app=app),
        base_url="http://testserver",
    ) as client:
        # 1. Login Admin A
        await login_with_totp(client, "admin_a_rep@tenant.com", PASSWORD)

        # 2. Case -> Doc -> Worker -> Features -> Search -> Candidates -> Comparison
        res_case = await client.post(
            f"/api/v1/organizations/{ORG_A}/cases",
            json={
                "case_number": "2026-REP-001",
                "title": "大模型注意力量化方法",
                "technical_field": "人工智能",
            },
        )
        assert res_case.status_code == 201
        case_id = res_case.json()["case"]["id"]

        doc = Document()
        doc.add_heading("技术领域", level=1)
        doc.add_paragraph("本申请涉及一种面向自注意力机制的大模型量化方法。")
        doc.add_heading("权利要求书", level=1)
        doc.add_paragraph("1. 一种面向自注意力的量化方法，其特征在于包括：针对Q、K、V矩阵进行非线性低位宽定点化。")
        buf = io.BytesIO()
        doc.save(buf)

        import base64
        await client.post(
            f"/api/v1/organizations/{ORG_A}/cases/{case_id}/documents",
            json={
                "filename": "claims.docx",
                "content_base64": base64.b64encode(buf.getvalue()).decode("ascii"),
                "content_type": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            },
        )

        worker_engine = create_engine(settings.worker_database_url)
        worker_session_factory = create_worker_session_factory(worker_engine)
        storage = ObjectStorageClient()
        worker = ParseTaskWorker(worker_session_factory, storage, clock=clock)
        await worker.run_once()
        await worker_engine.dispose()

        case_detail = (await client.get(f"/api/v1/organizations/{ORG_A}/cases/{case_id}")).json()
        doc_ver_id = case_detail["document_version"]["id"]
        await client.post(
            f"/api/v1/organizations/{ORG_A}/cases/{case_id}/document-versions/{doc_ver_id}/confirm"
        )

        res_feats = await client.post(
            f"/api/v1/organizations/{ORG_A}/cases/{case_id}/features/extract"
        )
        f_ver_id = res_feats.json()["version"]["id"]
        await client.post(
            f"/api/v1/organizations/{ORG_A}/cases/{case_id}/features/versions/{f_ver_id}/confirm"
        )

        await client.post(
            f"/api/v1/organizations/{ORG_A}/cases/{case_id}/search/strategies/generate"
        )
        cnipr_csv = """公开号,发明名称,摘要,公开日,申请人,IPC分类号
CN118888888A,自注意力量化系统,对QKV矩阵进行低秩分解与量化,2024-04-01,创新院,G06N 3/08
"""
        await client.post(
            f"/api/v1/organizations/{ORG_A}/cases/{case_id}/search/candidates/import-cnipr",
            json={"raw_content": cnipr_csv},
        )

        cands = (await client.get(f"/api/v1/organizations/{ORG_A}/cases/{case_id}/search/candidates")).json()["items"]
        await client.post(
            f"/api/v1/organizations/{ORG_A}/cases/{case_id}/search/candidates/{cands[0]['id']}/triage",
            json={"triage_status": "included"},
        )

        res_matrix = await client.post(
            f"/api/v1/organizations/{ORG_A}/cases/{case_id}/comparisons/generate"
        )
        matrix_id = res_matrix.json()["matrix"]["id"]
        await client.post(
            f"/api/v1/organizations/{ORG_A}/cases/{case_id}/comparisons/{matrix_id}/confirm"
        )

        # 3. Seal Evidence Snapshot
        res_seal = await client.post(
            f"/api/v1/organizations/{ORG_A}/cases/{case_id}/evidence/seal"
        )
        assert res_seal.status_code == 201
        seal_data = res_seal.json()
        assert seal_data["snapshot"]["status"] == "sealed"
        assert len(seal_data["snapshot"]["root_sha256"]) == 64
        assert "SNAP-" in seal_data["snapshot"]["snapshot_number"]
        assert seal_data["report"] is not None
        assert "# 专利证据分析与法律评估报告" in seal_data["report"]["content"]

        # Verify case status updated to completed
        case_end = (await client.get(f"/api/v1/organizations/{ORG_A}/cases/{case_id}")).json()
        assert case_end["status"] == "completed"

        # 4. Get active report
        res_rep = await client.get(
            f"/api/v1/organizations/{ORG_A}/cases/{case_id}/reports/active"
        )
        assert res_rep.status_code == 200
        assert res_rep.json()["report"]["title"] is not None

        # 5. Cross-Tenant Isolation: Admin B in Org B cannot read Org A's reports or snapshots
        await login_with_totp(client, "admin_b_rep@tenant.com", PASSWORD)
        res_cross_snap = await client.get(
            f"/api/v1/organizations/{ORG_B}/cases/{case_id}/evidence/active"
        )
        assert res_cross_snap.status_code == 404
        res_cross_rep = await client.get(
            f"/api/v1/organizations/{ORG_B}/cases/{case_id}/reports/active"
        )
        assert res_cross_rep.status_code == 404
