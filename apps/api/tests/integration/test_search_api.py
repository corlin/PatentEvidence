import io
import json
from datetime import UTC, datetime, timedelta
from uuid import UUID

import psycopg
import pytest
from argon2 import PasswordHasher
from cryptography.fernet import Fernet
from docx import Document
from httpx import ASGITransport, AsyncClient

from adapters.object_storage.client import ObjectStorageClient
from apps.worker.src.patent_evidence_worker.parse_worker import ParseTaskWorker
from patent_evidence_api.core.database import (
    create_engine,
    create_worker_session_factory,
)
from patent_evidence_api.main import create_app
from test_support import api_settings, login_with_totp, postgres_url

ORG_A = UUID("80000000-0000-4000-8000-000000000001")
ORG_B = UUID("80000000-0000-4000-8000-000000000002")
ADMIN_A = UUID("81000000-0000-4000-8000-000000000001")
ADMIN_B = UUID("81000000-0000-4000-8000-000000000002")
MEMBER_A = UUID("82000000-0000-4000-8000-000000000001")
MEMBER_B = UUID("82000000-0000-4000-8000-000000000002")
PASSWORD = "Correct horse battery staple 42"
MFA_KEY = Fernet.generate_key().decode()


class Clock:
    def __init__(self) -> None:
        self.value = datetime.now(UTC).replace(microsecond=0)

    def __call__(self) -> datetime:
        return self.value


def _seed_org_and_identities(migration_url: str, now: datetime) -> None:
    password_hash = PasswordHasher().hash(PASSWORD)
    with psycopg.connect(migration_url, autocommit=True) as conn:
        with conn.cursor() as cur:
            cur.execute(
                """INSERT INTO global_identities
                (id, email_normalized, status, security_version, created_at, updated_at)
                VALUES
                (%s, 'admin_a_search@tenant.com', 'active', 1, %s, %s),
                (%s, 'admin_b_search@tenant.com', 'active', 1, %s, %s)
                ON CONFLICT (id) DO NOTHING""",
                (ADMIN_A, now, now, ADMIN_B, now, now),
            )
            cur.execute(
                """INSERT INTO credentials
                (id, identity_id, credential_type, secret_hash, status, created_at, updated_at)
                VALUES
                (gen_random_uuid(), %s, 'password', %s, 'active', %s, %s),
                (gen_random_uuid(), %s, 'password', %s, 'active', %s, %s)
                ON CONFLICT DO NOTHING""",
                (ADMIN_A, password_hash, now, now, ADMIN_B, password_hash, now, now),
            )
            cur.execute(
                """INSERT INTO organizations
                (id, slug, display_name, persisted_status, created_at, updated_at)
                VALUES
                (%s, 'search-org-a', '检索机构A', 'active', %s, %s),
                (%s, 'search-org-b', '检索机构B', 'active', %s, %s)
                ON CONFLICT (id) DO NOTHING""",
                (ORG_A, now, now, ORG_B, now, now),
            )
            cur.execute(
                """INSERT INTO organization_plan_quotas
                (id, organization_id, plan_key, monthly_case_allowance, current_period_start, current_period_end, quota_status, created_at, updated_at)
                VALUES
                (gen_random_uuid(), %s, 'standard', 50, %s, %s, 'active', %s, %s),
                (gen_random_uuid(), %s, 'standard', 50, %s, %s, 'active', %s, %s)
                ON CONFLICT DO NOTHING""",
                (
                    ORG_A,
                    now - timedelta(days=1),
                    now + timedelta(days=29),
                    now,
                    now,
                    ORG_B,
                    now - timedelta(days=1),
                    now + timedelta(days=29),
                    now,
                    now,
                ),
            )
            cur.execute(
                """INSERT INTO organization_memberships
                (id, organization_id, identity_id, role, status, created_at, updated_at)
                VALUES
                (%s, %s, %s, 'organization_admin', 'active', %s, %s),
                (%s, %s, %s, 'organization_admin', 'active', %s, %s)
                ON CONFLICT (id) DO NOTHING""",
                (MEMBER_A, ORG_A, ADMIN_A, now, now, MEMBER_B, ORG_B, ADMIN_B, now, now),
            )


@pytest.mark.asyncio
async def test_search_planning_and_candidate_triage_lifecycle() -> None:
    migration_url = postgres_url("PE_MIGRATION_DATABASE_URL", "patent_evidence_migration")
    clock = Clock()
    _seed_org_and_identities(migration_url, clock())

    settings = api_settings(mfa_encryption_key=MFA_KEY)
    app = create_app(settings=settings, clock=clock)

    async with AsyncClient(
        transport=ASGITransport(app=app),
        base_url="http://testserver",
    ) as client:
        # 1. Login Admin A
        await login_with_totp(client, "admin_a_search@tenant.com", PASSWORD)

        # 2. Create Case & Document & Parse
        res_case = await client.post(
            f"/api/v1/organizations/{ORG_A}/cases",
            json={
                "case_number": "2026-SRCH-001",
                "title": "大语言模型混合精度量化推理加速方法",
                "technical_field": "人工智能与大模型计算",
            },
        )
        assert res_case.status_code == 201
        case_id = res_case.json()["case"]["id"]

        doc = Document()
        doc.add_heading("技术领域", level=1)
        doc.add_paragraph("本申请涉及一种面向低位宽的大语言模型量化推理加速系统。")
        doc.add_heading("发明内容", level=1)
        doc.add_paragraph("本发明通过获取神经网络权重并执行奇异值分解与稀疏量化查找表构建。")
        buf = io.BytesIO()
        doc.save(buf)

        import base64
        await client.post(
            f"/api/v1/organizations/{ORG_A}/cases/{case_id}/documents",
            json={
                "filename": "disclosure.docx",
                "content_base64": base64.b64encode(buf.getvalue()).decode("ascii"),
                "content_type": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            },
        )

        worker_engine = create_engine(settings.worker_database_url)
        worker_session_factory = create_worker_session_factory(worker_engine)
        storage = ObjectStorageClient()
        worker = ParseTaskWorker(worker_session_factory, storage, clock=clock)
        await worker.run_once()
        await worker_engine.dispose()

        # Confirm Document Version
        case_detail = (await client.get(f"/api/v1/organizations/{ORG_A}/cases/{case_id}")).json()
        doc_ver_id = case_detail["document_version"]["id"]
        await client.post(
            f"/api/v1/organizations/{ORG_A}/cases/{case_id}/document-versions/{doc_ver_id}/confirm"
        )

        # Extract & Confirm Features
        res_feats = await client.post(
            f"/api/v1/organizations/{ORG_A}/cases/{case_id}/features/extract"
        )
        f_ver_id = res_feats.json()["version"]["id"]
        await client.post(
            f"/api/v1/organizations/{ORG_A}/cases/{case_id}/features/versions/{f_ver_id}/confirm"
        )

        # 3. Generate Search Strategy
        res_strat = await client.post(
            f"/api/v1/organizations/{ORG_A}/cases/{case_id}/search/strategies/generate"
        )
        assert res_strat.status_code == 201
        strat_data = res_strat.json()
        strat_id = strat_data["id"]
        assert "AND" in strat_data["boolean_query_cnipr"]
        assert len(strat_data["ipc_classes"]) > 0

        # 4. Get CNIPR Handoff Package
        res_handoff = await client.get(
            f"/api/v1/organizations/{ORG_A}/cases/{case_id}/search/strategies/{strat_id}/handoff"
        )
        assert res_handoff.status_code == 200
        handoff_data = res_handoff.json()
        assert "CNIPR 官方专利检索人工交接规范包" in handoff_data["markdown"]
        assert handoff_data["json"]["case"]["case_number"] == "2026-SRCH-001"

        # 5. Execute Public Search (Google Patents)
        res_pub_search = await client.post(
            f"/api/v1/organizations/{ORG_A}/cases/{case_id}/search/jobs/execute-public",
            json={"source_type": "google_patents"},
        )
        assert res_pub_search.status_code == 201
        assert res_pub_search.json()["results_count"] > 0

        # 6. Import CNIPR Results (CSV with duplicate check)
        cnipr_csv = """公开号,发明名称,摘要,公开日,申请人,IPC分类号
CN117283912A,一种基于量化的多尺度模型量化加速装置及系统,通过奇异值分解降低张量冗余,2024-03-15,前沿智能,G06N 3/08
CN119999888A,另一种大模型稀疏化装置,针对大模型注意力层进行量化,2024-04-01,创新研究院,G06N 3/08
"""
        res_import = await client.post(
            f"/api/v1/organizations/{ORG_A}/cases/{case_id}/search/candidates/import-cnipr",
            json={"raw_content": cnipr_csv},
        )
        assert res_import.status_code == 201
        assert res_import.json()["imported_count"] == 2

        # 7. List Candidates (sorted by relevance_score)
        res_cands = await client.get(
            f"/api/v1/organizations/{ORG_A}/cases/{case_id}/search/candidates"
        )
        assert res_cands.status_code == 200
        cands = res_cands.json()["items"]
        assert len(cands) >= 4
        # Verify deduplication of CN117283912A
        cn117_matches = [c for c in cands if c["publication_number_normalized"] == "CN117283912A"]
        assert len(cn117_matches) == 1
        cand_1 = cands[0]
        cand_2 = cands[1]

        # 8. Triage Candidate 1 (Include) and Candidate 2 (Exclude)
        res_triage_inc = await client.post(
            f"/api/v1/organizations/{ORG_A}/cases/{case_id}/search/candidates/{cand_1['id']}/triage",
            json={"triage_status": "included", "notes": "该专利技术方案与本申请高度吻合"},
        )
        assert res_triage_inc.status_code == 200
        assert res_triage_inc.json()["triage_status"] == "included"

        res_triage_exc = await client.post(
            f"/api/v1/organizations/{ORG_A}/cases/{case_id}/search/candidates/{cand_2['id']}/triage",
            json={
                "triage_status": "excluded",
                "exclusion_reason": "缺乏关键特征F2",
                "notes": "未涉及矩阵奇异值分解步骤",
            },
        )
        assert res_triage_exc.status_code == 200
        assert res_triage_exc.json()["triage_status"] == "excluded"
        assert res_triage_exc.json()["exclusion_reason"] == "缺乏关键特征F2"

        # 9. Filter Candidates by status
        res_filtered = await client.get(
            f"/api/v1/organizations/{ORG_A}/cases/{case_id}/search/candidates?triage_status=included"
        )
        assert res_filtered.status_code == 200
        assert len(res_filtered.json()["items"]) == 1
        assert res_filtered.json()["items"][0]["id"] == cand_1["id"]

        # 10. Cross-Tenant Isolation: Admin B cannot access Org A's search strategy or candidates
        await login_with_totp(client, "admin_b_search@tenant.com", PASSWORD)
        res_cross = await client.get(
            f"/api/v1/organizations/{ORG_B}/cases/{case_id}/search/strategies/active"
        )
        assert res_cross.status_code == 404
