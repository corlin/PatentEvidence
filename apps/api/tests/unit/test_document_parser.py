from __future__ import annotations

import io
from docx import Document
from pypdf import PdfWriter
from modules.cases.parser import DocumentParser


def test_parse_plain_text() -> None:
    parser = DocumentParser()
    text = "第一段文本。\n第二段文本，描述技术方案。\n第三段文本。"
    result = parser.parse_bytes(text.encode("utf-8"), filename="test.txt")

    assert len(result.paragraphs) == 3
    assert result.paragraphs[0]["id"] == "p1"
    assert result.paragraphs[0]["text"] == "第一段文本。"
    assert result.paragraphs[1]["id"] == "p2"
    assert result.paragraphs[1]["text"] == "第二段文本，描述技术方案。"
    assert not result.is_scanned
    assert len(result.sha256) == 64


def test_parse_docx_with_headings_and_tables() -> None:
    parser = DocumentParser()

    # Generate an in-memory docx
    doc = Document()
    doc.add_heading("技术领域", level=1)
    doc.add_paragraph("本发明属于人工智能与大模型量化技术领域。")
    doc.add_heading("发明内容", level=1)
    doc.add_paragraph("为了解决现有技术中显存占用大的问题，本发明提出了一种混合精度量化方法。")

    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "参数名称"
    table.cell(0, 1).text = "位宽"
    table.cell(1, 0).text = "权重 W"
    table.cell(1, 1).text = "4-bit"

    buffer = io.BytesIO()
    doc.save(buffer)
    docx_bytes = buffer.getvalue()

    result = parser.parse_bytes(docx_bytes, filename="invention.docx")

    assert len(result.paragraphs) >= 4
    assert result.paragraphs[0]["section"] == "技术领域"
    assert "本发明属于" in result.paragraphs[1]["text"]
    assert result.paragraphs[2]["section"] == "发明内容"
    assert any("表格 1" in p["section"] for p in result.paragraphs)
    assert not result.is_scanned


def test_parse_pdf() -> None:
    parser = DocumentParser()

    # Create empty PDF
    writer = PdfWriter()
    writer.add_blank_page(width=72, height=72)
    buffer = io.BytesIO()
    writer.write(buffer)
    pdf_bytes = buffer.getvalue()

    result = parser.parse_bytes(pdf_bytes, filename="sample.pdf")
    # Blank PDF has < 50 chars, so flagged as scanned
    assert result.is_scanned
