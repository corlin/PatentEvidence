"""Unit tests for the Markdown → DOCX report exporter."""

from __future__ import annotations

import io
import zipfile
from datetime import datetime, timezone

from docx import Document

from modules.reports.docx_export import (
    ReportProvenance,
    markdown_to_docx_bytes,
    provenance_footer_text,
)

# 转换器不新增任何措辞，肯定式结论措辞不允许出现在导出产物里。
BANNED_PHRASES = ["良好授权前景", "创新高度", "全案风险评级"]


def _reopen(data: bytes) -> Document:
    return Document(io.BytesIO(data))


def test_converts_headings_paragraphs_and_bullets() -> None:
    md = "\n".join(
        [
            "# 专利证据分析与法律评估报告",
            "",
            "这是正文段落，含 **加粗** 与 `代码` 混排。",
            "",
            "## 1. 特征分解",
            "",
            "- 第一条要点",
            "- 第二条要点",
            "",
            "> 免责声明：本报告不构成专利性结论。",
            "",
            "---",
            "",
            "### 小节",
        ]
    )
    doc = _reopen(markdown_to_docx_bytes(md, title="测试报告"))

    headings = [p for p in doc.paragraphs if p.style.name.startswith("Heading") or p.style.name == "Title"]
    heading_texts = [p.text for p in headings]
    assert any("测试报告" in t for t in heading_texts)
    assert any("专利证据分析与法律评估报告" in t for t in heading_texts)
    assert any("特征分解" in t for t in heading_texts)

    bullets = [p for p in doc.paragraphs if p.style.name == "List Bullet"]
    assert [b.text for b in bullets] == ["第一条要点", "第二条要点"]

    quotes = [p for p in doc.paragraphs if "Quote" in p.style.name]
    assert any("不构成专利性结论" in q.text for q in quotes)

    # 加粗行内标记被正确转换而不是原样输出
    body = "\n".join(p.text for p in doc.paragraphs)
    assert "**" not in body
    assert "加粗" in body


def test_converts_pipe_tables_with_separator_rows() -> None:
    md = "\n".join(
        [
            "## 证据完备度",
            "",
            "| 来源覆盖率 | 已核验引文 | 阻断结论 |",
            "| :--- | :--- | :--- |",
            "| 0.8 | 3/4 | 否 |",
            "| 0.6 | 1/2 | 是 |",
            "",
            "表格后的段落。",
        ]
    )
    doc = _reopen(markdown_to_docx_bytes(md))

    assert len(doc.tables) == 1
    table = doc.tables[0]
    assert len(table.rows) == 3  # 表头 + 两行数据
    assert len(table.columns) == 3
    assert table.rows[0].cells[0].text == "来源覆盖率"
    assert table.rows[1].cells[2].text == "否"
    assert table.rows[2].cells[2].text == "是"
    # 分隔行不得泄漏成数据行
    assert all(":---" not in cell.text for row in table.rows for cell in row.cells)
    # 表头加粗
    assert all(run.bold for run in table.rows[0].cells[0].paragraphs[0].runs)


def test_handles_edge_cases_without_crashing() -> None:
    # 空内容、孤立竖线、残缺表格、行尾表格都不应抛错
    for md in [
        "",
        "\n\n\n",
        "| 只有一行表头 |",
        "| a | b |\n| :--- | :--- |",
        "普通段落 | 含竖线但不是表格",
        "- \n- 空要点",
    ]:
        data = markdown_to_docx_bytes(md)
        assert data.startswith(b"PK")  # 合法 docx（zip 魔数）


def test_export_adds_no_wording_of_its_own() -> None:
    md = "# 报告\n\n内容含否定式声明：不构成良好授权前景意见。\n"
    data = markdown_to_docx_bytes(md)
    doc = _reopen(data)
    body = "\n".join(p.text for p in doc.paragraphs)
    # 否定式声明照实保留（不能误删），但转换器自身不引入肯定式措辞
    assert "不构成良好授权前景意见" in body
    for phrase in ["具备良好授权前景", "创新高度", "全案风险评级"]:
        assert phrase not in body


def test_converts_real_generator_output_end_to_end() -> None:
    """用真实 MarkdownReportGenerator 产出的报告做端到端转换。"""
    from modules.reports.generator import MarkdownReportGenerator

    payload = {
        "case": {"title": "大模型量化", "case_number": "2026-CASE-001"},
        "document": {"filename": "交底书.docx", "version_number": 1, "file_sha256": "cd" * 32},
        "features": {
            "items": [
                {
                    "feature_code": "F1",
                    "feature_type": "preamble",
                    "source_paragraph_id": "[0009]",
                    "feature_statement": "一种量化方法",
                }
            ]
        },
        "search": {"boolean_query_cnipr": "TIAB=(量化 AND 模型)"},
        "candidates": [
            {
                "publication_number": "CN1A",
                "title": "对比文件一",
                "triage_status": "included",
            }
        ],
        "assessment": None,
    }
    md = MarkdownReportGenerator().generate(payload, "ab" * 32)
    data = markdown_to_docx_bytes(md, title="大模型量化 分析报告")
    doc = _reopen(data)

    body = "\n".join(p.text for p in doc.paragraphs)
    table_text = "\n".join(
        cell.text for table in doc.tables for row in table.rows for cell in row.cells
    )
    assert "大模型量化" in body or "大模型量化" in table_text
    for phrase in BANNED_PHRASES:
        assert phrase not in body
        assert phrase not in table_text


PROVENANCE = ReportProvenance(
    snapshot_number=7,
    root_sha256="ab" * 32,
    sealed_at=datetime(2026, 9, 1, 8, 30, tzinfo=timezone.utc),
)


def test_footer_carries_snapshot_version_and_root_hash() -> None:
    data = markdown_to_docx_bytes("# 报告\n\n正文", provenance=PROVENANCE)
    document = Document(io.BytesIO(data))
    footer = document.sections[0].footer.paragraphs[0].text
    assert footer == provenance_footer_text(PROVENANCE)
    assert "#7" in footer and "ab" * 32 in footer
    assert "2026-09-01T08:30:00Z" in footer
    # 页脚只在页脚，不混入正文
    assert all("快照根校验值" not in p.text for p in document.paragraphs)


def test_same_input_yields_identical_bytes_regardless_of_wall_clock() -> None:
    first = markdown_to_docx_bytes("# 报告", provenance=PROVENANCE)
    second = markdown_to_docx_bytes("# 报告", provenance=PROVENANCE)
    assert first == second
    stamps = {info.date_time for info in zipfile.ZipFile(io.BytesIO(first)).infolist()}
    assert stamps == {(2026, 9, 1, 8, 30, 0)}
    core = Document(io.BytesIO(first)).core_properties
    assert core.created == datetime(2026, 9, 1, 8, 30, tzinfo=timezone.utc)
    assert core.author == "PatentEvidence"


def test_different_snapshot_changes_the_file_bytes() -> None:
    other = ReportProvenance(8, "cd" * 32, PROVENANCE.sealed_at)
    assert markdown_to_docx_bytes("# 报告", provenance=PROVENANCE) != (
        markdown_to_docx_bytes("# 报告", provenance=other)
    )
