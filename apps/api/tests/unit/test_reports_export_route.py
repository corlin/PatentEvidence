"""Route tests for the DOCX/PDF exports of the active analysis report."""

from __future__ import annotations

import hashlib
import io
from contextlib import asynccontextmanager
from typing import Any, AsyncIterator
from uuid import UUID, uuid4

from docx import Document
from pypdf import PdfReader
from fastapi import FastAPI
from fastapi.testclient import TestClient

from patent_evidence_api.organization.access import OrganizationMutation
from patent_evidence_api.reports.api import create_reports_router
from patent_evidence_api.reports.services import EvidenceReportService

ORG = uuid4()
CASE = uuid4()

DOCX_MEDIA_TYPE = (
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
)

REPORT_MARKDOWN = "\n".join(
    [
        "# 专利证据分析与法律评估报告",
        "",
        "> **案件编号**：`2026-CASE-001`  ",
        "",
        "## 1. 技术特征分解",
        "",
        "| 特征编号 | 特征内容 |",
        "| :--- | :--- |",
        "| `F1` | 一种量化方法 |",
        "",
        "本报告不构成专利性结论。",
    ]
)


SNAPSHOT_ID = uuid4()
ROOT_SHA = "c0ffee" + "0" * 58


class FakeAccess:
    """替身 OrganizationAccess：记录 mutation 的审计动作与描述，不做真实鉴权。"""

    def __init__(self) -> None:
        self.mutations: list[dict[str, Any]] = []

    @asynccontextmanager
    async def mutation(
        self,
        request: Any,
        organization_id: UUID,
        *,
        action: str,
        target_type: str,
        target_id: UUID | None = None,
    ) -> AsyncIterator[OrganizationMutation]:
        operation = OrganizationMutation(
            session=object(),  # type: ignore[arg-type]
            principal=object(),  # type: ignore[arg-type]
            organization_id=organization_id,
            correlation_id="test",
            target_type=target_type,
            target_id=target_id,
        )
        record = {"action": action, "target_type": target_type, "result": "denied"}
        self.mutations.append(record)
        yield operation
        record.update(
            result="allowed",
            target_id=operation.target_id,
            safe_summary=operation.safe_summary,
        )


def _snapshot_payload() -> dict[str, Any]:
    return {
        "snapshot": {
            "id": str(SNAPSHOT_ID),
            "snapshot_number": 4,
            "root_sha256": ROOT_SHA,
            "sealed_at": "2026-09-20T03:04:05+00:00",
        },
        "report": {
            "id": str(uuid4()),
            "title": "大模型量化 分析报告",
            "content": REPORT_MARKDOWN,
        },
    }


class FakeReportService(EvidenceReportService):
    """返回罐头快照数据；不触碰数据库。"""

    def __init__(self, snapshot: dict[str, Any] | None) -> None:
        super().__init__()
        self._snapshot = snapshot

    async def get_active_snapshot(
        self, session: Any, organization_id: UUID, case_id: UUID
    ) -> dict[str, Any] | None:
        return self._snapshot


def _client(service: FakeReportService) -> tuple[TestClient, FakeAccess]:
    access = FakeAccess()
    app = FastAPI()
    app.include_router(create_reports_router(access, service))  # type: ignore[arg-type]
    return TestClient(app, raise_server_exceptions=False), access


def test_docx_export_returns_a_valid_word_attachment() -> None:
    service = FakeReportService(_snapshot_payload())
    client, access = _client(service)

    response = client.get(
        f"/api/v1/organizations/{ORG}/cases/{CASE}/reports/active/export.docx"
    )

    assert response.status_code == 200
    assert response.headers["content-type"].startswith(DOCX_MEDIA_TYPE)
    disposition = response.headers["content-disposition"]
    assert "attachment" in disposition
    assert "analysis-report.docx" in disposition

    # 产物是合法 docx，内容来自封存报告本身（含免责声明、无新增措辞）
    doc = Document(io.BytesIO(response.content))
    body = "\n".join(p.text for p in doc.paragraphs)
    table_text = "\n".join(
        cell.text for table in doc.tables for row in table.rows for cell in row.cells
    )
    assert "专利证据分析与法律评估报告" in body
    assert "不构成专利性结论" in body
    assert "一种量化方法" in table_text
    for banned in ("良好授权前景", "创新高度", "全案风险评级"):
        assert banned not in body
        assert banned not in table_text
    footer = doc.sections[0].footer.paragraphs[0].text
    assert "证据快照版本 #4" in footer
    assert ROOT_SHA in footer

    # 文件哈希经响应头公开，并写入审计记录
    file_sha = hashlib.sha256(response.content).hexdigest()
    assert response.headers["x-content-sha256"] == file_sha
    assert response.headers["x-snapshot-root-sha256"] == ROOT_SHA
    assert access.mutations == [
        {
            "action": "case.report.export_docx",
            "target_type": "evidence_snapshot",
            "result": "allowed",
            "target_id": SNAPSHOT_ID,
            "safe_summary": f"Exported report DOCX for snapshot #4 file_sha256={file_sha}",
        }
    ]


def test_docx_export_bytes_are_stable_for_the_same_snapshot() -> None:
    client, _ = _client(FakeReportService(_snapshot_payload()))
    url = f"/api/v1/organizations/{ORG}/cases/{CASE}/reports/active/export.docx"
    first, second = client.get(url), client.get(url)
    assert first.content == second.content
    assert first.headers["x-content-sha256"] == second.headers["x-content-sha256"]


def test_docx_export_is_404_when_no_report_exists() -> None:
    client, access = _client(FakeReportService(None))
    response = client.get(
        f"/api/v1/organizations/{ORG}/cases/{CASE}/reports/active/export.docx"
    )
    assert access.mutations[0]["result"] == "denied"
    assert response.status_code == 404
    assert response.json()["detail"] == "no_report_found"


def test_docx_export_is_404_when_snapshot_has_no_report() -> None:
    client, _ = _client(FakeReportService({"snapshot": _snapshot_payload()["snapshot"], "report": None}))
    response = client.get(
        f"/api/v1/organizations/{ORG}/cases/{CASE}/reports/active/export.docx"
    )
    assert response.status_code == 404
    assert response.json()["detail"] == "no_report_found"


def test_pdf_export_returns_a_hashed_audited_pdf() -> None:
    client, access = _client(FakeReportService(_snapshot_payload()))
    url = f"/api/v1/organizations/{ORG}/cases/{CASE}/reports/active/export.pdf"

    response = client.get(url)

    assert response.status_code == 200
    assert response.headers["content-type"] == "application/pdf"
    assert "analysis-report.pdf" in response.headers["content-disposition"]
    assert response.content.startswith(b"%PDF-")

    text = "".join(page.extract_text() for page in PdfReader(io.BytesIO(response.content)).pages)
    assert "一种量化方法" in text
    assert "不构成专利性结论" in text
    assert "证据快照版本 #4" in text and ROOT_SHA in text

    file_sha = hashlib.sha256(response.content).hexdigest()
    assert response.headers["x-content-sha256"] == file_sha
    assert access.mutations == [
        {
            "action": "case.report.export_pdf",
            "target_type": "evidence_snapshot",
            "result": "allowed",
            "target_id": SNAPSHOT_ID,
            "safe_summary": f"Exported report PDF for snapshot #4 file_sha256={file_sha}",
        }
    ]
    # 同一快照再次导出，字节与哈希不变
    assert client.get(url).content == response.content


def test_pdf_export_is_404_when_no_report_exists() -> None:
    client, access = _client(FakeReportService(None))
    response = client.get(
        f"/api/v1/organizations/{ORG}/cases/{CASE}/reports/active/export.pdf"
    )
    assert response.status_code == 404
    assert access.mutations[0]["result"] == "denied"
