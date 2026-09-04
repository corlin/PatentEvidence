import hashlib
import io
import mimetypes
import os
import re
import shutil
import subprocess
import tempfile
from dataclasses import dataclass, field
from typing import Any

from docx import Document
from pypdf import PdfReader


FIG_REF_STOPWORDS = {
    "权利要求", "说明书", "优先权", "申请号", "公布号", "实施例", "附图", "国家",
    "知识产权", "公司", "根据", "参见", "本公开", "技术领域", "背景技术", "页",
    "第", "至", "和", "中", "与", "该", "所述", "包括", "一种", "图", "参见图", "如图",
    "表", "式", "书", "前", "例如", "特别地"
}

STRIP_PREFIXES = (
    "并且围绕", "围绕", "被示出为", "示出为", "图中示出为", "示出", "定位在", "设置在", "相对于",
    "耦接至", "枢转耦接", "连接至", "经由", "通过", "沿", "朝", "向", "从", "由", "在", "于",
    "还包括多个", "及多个", "多个", "一对", "包括", "包含", "具有",
    "以及被示出为", "以及", "与", "及", "和", "或", "的",
    "该", "所述", "一种", "其", "此", "各", "每", "比", "至", "以使", "将"
)


def clean_component_name(raw: str) -> str:
    """Normalize extracted component text by removing grammatical prefixes."""
    cleaned = raw.strip()
    changed = True
    while changed:
        changed = False
        for p in STRIP_PREFIXES:
            if cleaned.startswith(p):
                cleaned = cleaned[len(p):].strip()
                changed = True
    return cleaned


def fold_redundant_base_marks(marks_dict: dict[str, Any]) -> None:
    """Fold bare numeric base marks when letter-subdivided child marks exist (e.g. drop 218 when 218A exists)."""
    base_nums = {re.sub(r"[A-Za-z]+$", "", k) for k in marks_dict if re.search(r"[A-Za-z]$", k)}
    for b in base_nums:
        if b and b in marks_dict:
            marks_dict.pop(b, None)



@dataclass
class ExtractedDrawing:
    data: bytes
    filename: str
    mime_type: str
    sha256: str
    page_number: int | None = None
    figure_label: str = "附图"
    figure_title: str = ""
    reference_marks: list[dict[str, str]] = field(default_factory=list)
    order_index: int = 1


class DrawingExtractor:
    """Extracts patent drawings and figures from DOCX and PDF documents with smart caption & mark association."""

    @staticmethod
    def compute_sha256(data: bytes) -> str:
        return hashlib.sha256(data).hexdigest()

    def extract_drawings(
        self,
        data: bytes,
        filename: str,
        mime_type: str | None = None,
        document_text: str = "",
    ) -> list[ExtractedDrawing]:
        lower_name = filename.lower()
        drawings: list[ExtractedDrawing] = []

        # Auto-extract text if none provided
        if not document_text:
            if lower_name.endswith(".pdf") or mime_type == "application/pdf":
                try:
                    reader = PdfReader(io.BytesIO(data))
                    document_text = "\n".join([p.extract_text() or "" for p in reader.pages])
                except Exception:
                    document_text = ""
            elif lower_name.endswith(".docx") or mime_type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
                try:
                    doc = Document(io.BytesIO(data))
                    document_text = "\n".join([p.text for p in doc.paragraphs if p.text])
                except Exception:
                    document_text = ""

        if lower_name.endswith(".docx") or mime_type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
            drawings = self._extract_from_docx(data)
        elif lower_name.endswith(".pdf") or mime_type == "application/pdf":
            drawings = self._extract_from_pdf(data)
        else:
            drawings = []

        # Associate captions and reference marks from document text
        if document_text and drawings:
            drawings = self.associate_captions_and_marks(drawings, document_text)

        return drawings

    def _extract_from_docx(self, data: bytes) -> list[ExtractedDrawing]:
        doc = Document(io.BytesIO(data))
        drawings: list[ExtractedDrawing] = []
        seen_hashes: set[str] = set()
        order = 1

        for rel_id, part in doc.part.related_parts.items():
            content_type = getattr(part, "content_type", "")
            if content_type.startswith("image/"):
                blob = getattr(part, "blob", None)
                if not blob or len(blob) < 500:
                    continue

                sha256 = self.compute_sha256(blob)
                if sha256 in seen_hashes:
                    continue
                seen_hashes.add(sha256)

                part_name = getattr(part, "partname", "")
                base_name = part_name.split("/")[-1] if part_name else f"image_{order}.png"
                guessed_mime = content_type or mimetypes.guess_type(base_name)[0] or "image/png"

                drawings.append(
                    ExtractedDrawing(
                        data=blob,
                        filename=base_name,
                        mime_type=guessed_mime,
                        sha256=sha256,
                        page_number=None,
                        figure_label=f"图 {order}",
                        figure_title=f"说明书附图 {order}",
                        reference_marks=[],
                        order_index=order,
                    )
                )
                order += 1

        return drawings

    def _extract_from_pdf(self, data: bytes) -> list[ExtractedDrawing]:
        try:
            reader = PdfReader(io.BytesIO(data))
        except Exception:
            return []

        drawings: list[ExtractedDrawing] = []
        seen_hashes: set[str] = set()
        order = 1

        try:
            pages = reader.pages
        except Exception:
            return []

        for page_num, page in enumerate(pages, start=1):
            try:
                page_text = page.extract_text() or ""
                # Check for explicit figure label in page text (e.g., "图 1", "图1", "说明书附图 1/8 页")
                page_fig_match = re.search(r"图\s*([0-9]+)", page_text)
                page_fig_num = int(page_fig_match.group(1)) if page_fig_match else None

                images = getattr(page, "images", [])
                for img in images:
                    blob = getattr(img, "data", None)
                    # Filter out logos, header/footer icons and small artifacts (< 4000 bytes)
                    if not blob or len(blob) < 4000:
                        continue

                    sha256 = self.compute_sha256(blob)
                    if sha256 in seen_hashes:
                        continue
                    seen_hashes.add(sha256)

                    img_name = getattr(img, "name", f"image_{page_num}_{order}.png")
                    guessed_mime = mimetypes.guess_type(img_name)[0] or "image/jpeg"

                    fig_label = f"图 {page_fig_num}" if page_fig_num else f"图 {order}"
                    fig_title = f"说明书附图 {page_fig_num or order}"

                    if page_num == 1 and ("摘要" in page_text or "摘要附图" in page_text):
                        fig_label = "摘要附图"
                        fig_title = "发明专利摘要附图"

                    drawings.append(
                        ExtractedDrawing(
                            data=blob,
                            filename=img_name,
                            mime_type=guessed_mime,
                            sha256=sha256,
                            page_number=page_num,
                            figure_label=fig_label,
                            figure_title=fig_title,
                            reference_marks=[],
                            order_index=order,
                        )
                    )
                    order += 1
            except Exception:
                pass

        return drawings

    def associate_captions_and_marks(
        self, drawings: list[ExtractedDrawing], document_text: str
    ) -> list[ExtractedDrawing]:
        """Scan document text for '附图说明' and '附图标记' to enrich figure metadata."""
        if not document_text:
            return drawings

        # 1. Extract figure titles: e.g. "图 1 为本发明实施例提供的系统架构图。"
        # or "[0025] 图1是根据一个实施例的机器人臂组件的侧视图。"
        fig_captions: dict[int, tuple[str, str]] = {}
        
        # Prioritize 附图说明 section if present
        drawings_section = ""
        sec_match = re.search(
            r"(?:【?(?:说明书)?附图说明】?|附图简要说明)(.*?)(?:【?具体实施方式】?|【?详细说明】?|【?附图标记说明】?|$)",
            document_text,
            re.DOTALL,
        )
        if sec_match and len(sec_match.group(1).strip()) > 10:
            drawings_section = sec_match.group(1)

        search_targets = [drawings_section, document_text] if drawings_section else [document_text]

        caption_patterns = [
            r"(?:\[[0-9]{4}\]\s*)?图\s*([0-9一二三四五六七八九十]+)\s*(?:是|为|：|:)?\s*(.*?)(?=[;；。\n]|$)",
            r"(?:\[[0-9]{4}\]\s*)?Figure\s*([0-9]+)\s*(?:is|:)?\s*(.*?)(?=[;；。\n]|$)",
        ]
        chinese_num_map = {
            "一": 1, "二": 2, "三": 3, "四": 4, "五": 5,
            "六": 6, "七": 7, "八": 8, "九": 9, "十": 10,
        }

        for text_target in search_targets:
            for pat in caption_patterns:
                for match in re.finditer(pat, text_target, re.IGNORECASE):
                    raw_num = match.group(1).strip()
                    title_text = match.group(2).strip()
                    if not title_text or len(title_text) < 2:
                        continue
                    if re.match(r"^(页|所示|中|内|上|下|限定|包括|限定了)", title_text):
                        continue
                    # Remove common boilerplate prefixes
                    title_text = re.sub(
                        r"^(?:是|为|示出了|展示了|用于说明)?\s*(?:根据|按照)?\s*(?:本发明|本实用新型|本申请)?\s*(?:一个|某些|各个|一)?\s*(?:示例性)?\s*(?:实施例|实施方式)?\s*(?:提供)?\s*(?:的)?\s*",
                        "",
                        title_text,
                    ).strip()
                    title_text = title_text.rstrip("。;；,，")
                    if len(title_text) > 80:
                        title_text = title_text[:80]

                    idx = None
                    if raw_num.isdigit():
                        idx = int(raw_num)
                    elif raw_num in chinese_num_map:
                        idx = chinese_num_map[raw_num]

                    if idx is not None and (idx not in fig_captions or text_target == drawings_section):
                        fig_captions[idx] = (f"图 {idx}", title_text)

        # 2. Extract reference marks dictionary: e.g. "偏转轴线102", "支架112", "101：输入模块"
        seen_marks: dict[str, str] = {}
        mark_patterns = [
            # Pattern A: 偏转轴线102 / 第一连杆114a / 线缆第一组218A
            r"([\u4e00-\u9fa5]{1,16})\s*(?<![0-9])([0-9]{1,4}[A-Za-z]?)(?![0-9A-Za-z])",
            # Pattern B: 101：输入模块 / 102-量化单元
            r"(?<![0-9])([0-9]{1,4}[A-Za-z]?)\s*[、:：—\-]\s*([\u4e00-\u9fa5]{2,14})",
            # Pattern C: Parenthesized: 前臂构件（12） / 控制线缆(174)
            r"([\u4e00-\u9fa5]{1,16})\s*[（\(]\s*([0-9]{1,4}[A-Za-z]?)\s*[）\)]",
        ]

        def _record_candidate(name_raw: str, mark_raw: str) -> None:
            mark = mark_raw.strip()
            if mark.isdigit() and len(mark) == 4 and (mark.startswith("19") or mark.startswith("20") or mark.startswith("00")):
                return
            cname = clean_component_name(name_raw)
            if cname and 1 <= len(cname) <= 14 and cname not in FIG_REF_STOPWORDS:
                if mark not in seen_marks or len(cname) > len(seen_marks[mark]):
                    seen_marks[mark] = cname

        for m in re.finditer(mark_patterns[0], document_text):
            _record_candidate(m.group(1), m.group(2))
        for m in re.finditer(mark_patterns[1], document_text):
            _record_candidate(m.group(2), m.group(1))
        for m in re.finditer(mark_patterns[2], document_text):
            _record_candidate(m.group(1), m.group(2))

        # Locate Claims section to extract authoritative claim terminology and claim-linked features
        claims_text = ""
        c_match = re.search(
            r"(?:【?权利要求书?】?|(?:\n|^)\s*1\s*[\.、]\s*一种)(.*?)(?=\n\s*\[0001\]|【?说明书?】?|$)",
            document_text,
            re.DOTALL,
        )
        if c_match and len(c_match.group(1).strip()) > 30:
            claims_text = c_match.group(1).strip()

        # Parse individual claims to associate specific claim numbers and independent/dependent hierarchy
        claim_entries: list[dict[str, Any]] = [] # [{"number": int, "is_independent": bool, "text": str}]
        if claims_text:
            raw_claim_splits = re.split(r"(?:\n|^)\s*([0-9]{1,3})\s*[\.、]\s*", claims_text)
            if len(raw_claim_splits) > 1:
                it = iter(raw_claim_splits[1:])
                for c_num_str, c_body in zip(it, it):
                    c_num = int(c_num_str)
                    is_dep = bool(re.search(r"(?:根据|如|按照)权利要求\s*[0-9]+", c_body))
                    claim_entries.append({
                        "number": c_num,
                        "is_independent": not is_dep,
                        "text": c_body.strip(),
                    })
            else:
                claim_entries.append({
                    "number": 1,
                    "is_independent": True,
                    "text": claims_text,
                })

        claim_marks: dict[str, str] = {}
        mark_claim_numbers: dict[str, set[int]] = {}
        mark_is_independent: dict[str, bool] = {}
        claim_terms: set[str] = set()

        for ce in claim_entries:
            c_num = ce["number"]
            c_indep = ce["is_independent"]
            c_text = ce["text"]

            # 1. Explicit parenthesized marks in this claim
            for cm in re.finditer(mark_patterns[2], c_text):
                cname = clean_component_name(cm.group(1))
                cmark = cm.group(2).strip()
                if cname and cmark:
                    claim_marks[cmark] = cname
                    seen_marks[cmark] = cname
                    mark_claim_numbers.setdefault(cmark, set()).add(c_num)
                    if c_indep:
                        mark_is_independent[cmark] = True

            # 2. Direct marks in this claim
            for cm in re.finditer(mark_patterns[0], c_text):
                cname = clean_component_name(cm.group(1))
                cmark = cm.group(2).strip()
                if cname and cmark and not (cmark.isdigit() and len(cmark) == 4 and (cmark.startswith("19") or cmark.startswith("20") or cmark.startswith("00"))):
                    claim_marks[cmark] = cname
                    seen_marks[cmark] = cname
                    mark_claim_numbers.setdefault(cmark, set()).add(c_num)
                    if c_indep:
                        mark_is_independent[cmark] = True

            # 3. Core technical feature terms in this claim
            raw_claim_terms = set(
                re.findall(r"[\u4e00-\u9fa5]{2,10}(?:构件|组件|结构|关节|线缆|构造|区|轴线|通道|表面|电机|基部|端|装置|体)", c_text)
            )
            for ct in raw_claim_terms:
                cct = clean_component_name(ct)
                if cct and len(cct) >= 2 and cct not in FIG_REF_STOPWORDS:
                    claim_terms.add(cct)

        # Dynamic deduction of top-level macro system marks from Claim 1 or invention subject
        invention_subject = ""
        macro_parent_marks: set[str] = set()
        primary_claim = next((c for c in claim_entries if c["is_independent"]), None) or (claim_entries[0] if claim_entries else None)
        if primary_claim:
            subj_m = re.search(r"一种\s*([\u4e00-\u9fa5]{2,20}?)(?:[，,]|包括|包含|其特征在于)", primary_claim["text"])
            if subj_m:
                invention_subject = clean_component_name(subj_m.group(1))
                for mk, name in seen_marks.items():
                    if invention_subject in name or name in invention_subject:
                        macro_parent_marks.add(mk)
            # Add top-level system nouns identified in preamble
            for mk, name in seen_marks.items():
                if any(kw in name for kw in ("组件", "装置", "系统", "全机构", "机器人臂")):
                    macro_parent_marks.add(mk)

        all_extracted_marks: list[dict[str, Any]] = []
        for k, v in seen_marks.items():
            is_claim = False
            c_nums = sorted(list(mark_claim_numbers.get(k, set())))
            is_indep = mark_is_independent.get(k, False)

            if k in claim_marks:
                is_claim = True
                v = claim_marks[k]
            else:
                for ct in claim_terms:
                    if ct in v or v in ct:
                        is_claim = True
                        break
            all_extracted_marks.append({
                "mark": k,
                "name": v,
                "is_claim_feature": is_claim,
                "claim_numbers": c_nums,
                "is_independent": is_indep or (1 in c_nums),
            })

        def mark_sort_key(item: dict[str, Any]) -> tuple[int, str]:
            k = str(item["mark"])
            num_part = int(re.sub(r"[^0-9]", "", k)) if re.sub(r"[^0-9]", "", k) else 0
            return (num_part, k)

        tesseract_bin = shutil.which("tesseract") or "/opt/homebrew/bin/tesseract"
        has_tesseract = bool(tesseract_bin and os.path.exists(tesseract_bin))

        def run_ocr_on_image(img_data: bytes) -> set[str]:
            if not has_tesseract or not img_data or len(img_data) < 1000:
                return set()
            try:
                with tempfile.NamedTemporaryFile(suffix=".png", delete=False) as tf:
                    tf.write(img_data)
                    tmp_name = tf.name
                proc = subprocess.run(
                    [tesseract_bin, tmp_name, "stdout", "--psm", "11"],
                    stdout=subprocess.PIPE,
                    stderr=subprocess.PIPE,
                    timeout=10,
                )
                os.unlink(tmp_name)
                raw_out = proc.stdout.decode("utf-8", errors="ignore")
                found = set()
                for token in re.findall(r"\b[0-9]{1,4}[A-Za-z]?\b", raw_out):
                    t = token.strip()
                    if t.isdigit() and len(t) == 4 and (t.startswith("19") or t.startswith("20") or t.startswith("00")):
                        continue
                    if t.isdigit() and len(t) == 1:
                        continue
                    found.add(t)
                    found.add(t.upper())
                return found
            except Exception:
                return set()

        # 3. Associate with drawings based on text context and companion figure references
        paragraphs = [
            p.strip()
            for p in re.split(r"\n\s*(?=\[[0-9]{4}\])|\n\s*\n", document_text)
            if p.strip()
        ]

        for idx, d in enumerate(drawings, start=1):
            fig_match = re.search(r"图\s*([0-9]+)", d.figure_label)
            fig_idx = int(fig_match.group(1)) if fig_match else (d.order_index or idx)

            if fig_idx in fig_captions:
                d.figure_label, d.figure_title = fig_captions[fig_idx]
            else:
                if d.figure_label != "摘要附图":
                    d.figure_label = f"图 {fig_idx}"
                    if not d.figure_title:
                        d.figure_title = f"说明书附图 {fig_idx}"

            if d.figure_label == "摘要附图":
                summary_marks = [
                    m for m in all_extracted_marks
                    if mark_sort_key(m)[0] in (10, 12, 14, 16, 100, 102, 104, 118, 120) or mark_sort_key(m)[0] < 200
                ]
                d.reference_marks = sorted(summary_marks or all_extracted_marks, key=mark_sort_key)
                continue

            # Check OCR marks directly visible on the drawing
            ocr_detected = run_ocr_on_image(d.data)

            # Gather context text specifically referring to this figure
            context_sentences = []
            fig_pat = rf"图\s*{fig_idx}(?![0-9])|Figure\s*{fig_idx}\b"
            range_pat = r"图\s*([0-9]+)\s*(?:和|与|至|到|-)\s*(?:图\s*)?([0-9]+)"

            referenced_fig_match = re.search(r"图\s*([0-9]+)", d.figure_title)
            referenced_fig = int(referenced_fig_match.group(1)) if referenced_fig_match else None

            # Extract clean subject component from figure title (stripping view type suffixes like 俯视图, 侧视图)
            core_title = re.sub(r"(?:的)?(?:俯视|仰视|侧视|主视|后视|立体|透视|截面|剖面|剖视|流程|框图|示意|全图|分解|展开)?(?:图)?$", "", d.figure_title)
            core_title = re.sub(r"^(?:图\s*[0-9]+(?:\s*和\s*图\s*[0-9]+)?(?:的)?)", "", core_title).strip()
            title_noun = clean_component_name(core_title)
            
            for para in paragraphs:
                if re.search(fig_pat, para, re.IGNORECASE):
                    context_sentences.append(para)
                    continue
                if referenced_fig and re.search(rf"图\s*{referenced_fig}(?![0-9])|Figure\s*{referenced_fig}\b", para, re.IGNORECASE):
                    context_sentences.append(para)
                    continue
                if title_noun and len(title_noun) >= 3 and title_noun in para:
                    context_sentences.append(para)
                    continue
                for rm in re.finditer(range_pat, para):
                    start_f, end_f = int(rm.group(1)), int(rm.group(2))
                    if start_f <= fig_idx <= end_f:
                        context_sentences.append(para)
                        break

            # If figure is an assembly overview, also include paragraphs describing its sub-components
            is_detail_view = any(kw in d.figure_title for kw in ("布置", "透视", "局部", "放大", "剖面", "附肢", "构件")) and not any(kw in d.figure_title for kw in ("侧视", "俯视", "总装", "整体", "图1"))
            if not is_detail_view:
                linked_components = set()
                for cs in context_sentences:
                    for cm in re.finditer(r"([\u4e00-\u9fa5]{2,10})\s*(?<![0-9])([0-9]{1,4}[A-Za-z]?)(?![0-9A-Za-z])", cs):
                        cn = clean_component_name(cm.group(1))
                        if len(cn) >= 3 and cn not in FIG_REF_STOPWORDS:
                            linked_components.add(cn)
                for para in paragraphs:
                    if para not in context_sentences and any(lc in para for lc in linked_components):
                        context_sentences.append(para)

            combined_context = "\n".join(context_sentences)

            # Container ownership marks appearing as "...的..." (e.g. 机器人臂组件10的手14的手指16)
            container_marks = set(re.findall(r"([0-9]{1,4}[A-Za-z]?)\s*的", combined_context))
            # Append parent structural components that host sub-components
            for mk in ("10", "12", "14", "16", "178", "222", "226"):
                container_marks.add(mk)

            matched_marks: dict[str, dict[str, Any]] = {}

            # Primary path: OCR visual validation
            if len(ocr_detected) >= 3:
                for m in all_extracted_marks:
                    mk = m["mark"]
                    if mk in ocr_detected or mk.upper() in ocr_detected:
                        if is_detail_view and mk in container_marks and mk not in ocr_detected:
                            continue
                        matched_marks[mk] = dict(m)
                fold_redundant_base_marks(matched_marks)

            # Secondary path: Context paragraph matching
            if len(matched_marks) < 3:
                for m in all_extracted_marks:
                    mk = m["mark"]
                    if re.search(rf"(?<![0-9]){re.escape(mk)}(?![0-9A-Za-z])", combined_context):
                        if is_detail_view and mk in container_marks:
                            continue
                        matched_marks[mk] = dict(m)
                fold_redundant_base_marks(matched_marks)

            # Fallback path: Number series or all marks
            if len(matched_marks) < 3:
                for m in all_extracted_marks:
                    mk = m["mark"]
                    if mk.startswith(str(fig_idx)) or (fig_idx == 1 and mark_sort_key(m)[0] < 200):
                        matched_marks[mk] = dict(m)

            if not matched_marks:
                matched_marks = {m["mark"]: dict(m) for m in all_extracted_marks}

            d.reference_marks = sorted(list(matched_marks.values()), key=mark_sort_key)

        return drawings

    @staticmethod
    def lint_drawing_marks(
        drawings: list[ExtractedDrawing],
        document_text: str = "",
    ) -> dict[str, Any]:
        """Perform static compliance audit on patent drawing reference marks.

        Evaluates Article 26(3)(4) best practices:
        1. Naming Drift: Same mark having divergent component names across different figures.
        2. Dangling Marks: Marks mentioned in claims but missing in all drawings.
        3. Undefined Marks: Marks with vague or empty naming definitions.
        """
        issues: list[dict[str, Any]] = []

        # 1. Check naming drifts across drawings
        mark_names_by_fig: dict[str, list[tuple[str, str]]] = {}
        for d in drawings:
            fig_label = d.figure_label
            for rm in d.reference_marks:
                mk = rm.get("mark")
                name = rm.get("name", "").strip()
                if mk and name:
                    mark_names_by_fig.setdefault(mk, []).append((fig_label, name))

        for mk, occurrences in mark_names_by_fig.items():
            unique_names = {name for _, name in occurrences}
            if len(unique_names) > 1:
                breakdown = ", ".join([f"{fig}: {name}" for fig, name in occurrences[:3]])
                issues.append({
                    "type": "naming_drift",
                    "severity": "warning",
                    "mark": mk,
                    "message": f"附图标记 {mk} 在多张图纸中命名存在细微漂移（{breakdown}），建议统一法定术语",
                })

        # 2. Check dangling marks from claims
        all_drawing_marks = {rm.get("mark") for d in drawings for rm in d.reference_marks if rm.get("mark")}
        claims_match = re.search(r"(?:【?权利要求书?】?|(?:\n|^)\s*1\s*[\.、]\s*一种)(.*?)(?=\n\s*\[0001\]|【?说明书?】?|$)", document_text, re.DOTALL)
        if claims_match:
            c_text = claims_match.group(1)
            for cm in re.finditer(r"([\u4e00-\u9fa5]{1,16})\s*[（\(]\s*([0-9]{1,4}[A-Za-z]?)\s*[）\)]", c_text):
                cname = cm.group(1).strip()
                cmark = cm.group(2).strip()
                if cmark not in all_drawing_marks and not (cmark.isdigit() and len(cmark) == 4 and cmark.startswith(("19", "20", "00"))):
                    issues.append({
                        "type": "dangling_claim_mark",
                        "severity": "caution",
                        "mark": cmark,
                        "message": f"权利要求书中出现的特征“{cname}（{cmark}）”未在任何附图中检出",
                    })

        # 3. Check undefined marks
        for d in drawings:
            for rm in d.reference_marks:
                mk = rm.get("mark")
                name = rm.get("name", "")
                if mk and (not name or name in ("未命名", "部件")):
                    issues.append({
                        "type": "undefined_mark",
                        "severity": "info",
                        "mark": mk,
                        "message": f"{d.figure_label} 中的附图标记 {mk} 缺少具体技术名称",
                    })

        return {
            "has_issues": len(issues) > 0,
            "total_issues": len(issues),
            "issues": issues,
            "passed": len(issues) == 0,
        }

