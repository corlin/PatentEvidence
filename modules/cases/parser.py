from __future__ import annotations

import hashlib
import io
import re
from dataclasses import asdict, dataclass, field
from typing import Any

from docx import Document
from pypdf import PdfReader


from modules.cases.drawing_extractor import DrawingExtractor, ExtractedDrawing


def split_sentences(text: str, min_length: int = 6) -> list[str]:
    """Split text into sentences using standard punctuation delimiters."""
    chunks = re.split(r"[；;。\n！!？?]", text)
    return [c.strip() for c in chunks if len(c.strip()) >= min_length]


@dataclass
class ParagraphBlock:
    id: str
    index: int
    section: str
    text: str
    offset_start: int
    offset_end: int


@dataclass
class ParsedDocument:
    parsed_text: str
    paragraphs: list[dict[str, Any]]
    sha256: str
    is_scanned: bool = False
    drawings: list[ExtractedDrawing] = field(default_factory=list)


class DocumentParser:
    """Extracts structured text and paragraph block objects from DOCX, PDF, or plain text."""

    def __init__(self) -> None:
        self._drawing_extractor = DrawingExtractor()

    @staticmethod
    def compute_sha256(data: bytes) -> str:
        return hashlib.sha256(data).hexdigest()

    def parse_bytes(
        self,
        data: bytes,
        filename: str,
        mime_type: str | None = None,
    ) -> ParsedDocument:
        sha256 = self.compute_sha256(data)
        lower_name = filename.lower()

        if lower_name.endswith(".docx") or mime_type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
            doc_result = self._parse_docx(data, sha256)
        elif lower_name.endswith(".pdf") or mime_type == "application/pdf":
            doc_result = self._parse_pdf(data, sha256)
        else:
            doc_result = self._parse_text(data.decode("utf-8", errors="replace"), sha256)

        # Extract drawings and enrich with document text
        drawings = self._drawing_extractor.extract_drawings(
            data, filename, mime_type, document_text=doc_result.parsed_text
        )
        doc_result.drawings = drawings
        return doc_result

    def _parse_docx(self, data: bytes, sha256: str) -> ParsedDocument:
        doc = Document(io.BytesIO(data))
        paragraphs: list[ParagraphBlock] = []
        full_text_chunks: list[str] = []
        current_offset = 0
        current_section = "正文"
        idx = 1

        for p in doc.paragraphs:
            raw_text = p.text.strip()
            if not raw_text:
                continue

            # Check if this paragraph is a heading
            style_name = p.style.name.lower() if p.style and p.style.name else ""
            if "heading" in style_name or "title" in style_name:
                current_section = raw_text[:64]

            offset_start = current_offset
            offset_end = offset_start + len(raw_text)

            paragraphs.append(
                ParagraphBlock(
                    id=f"p{idx}",
                    index=idx,
                    section=current_section,
                    text=raw_text,
                    offset_start=offset_start,
                    offset_end=offset_end,
                )
            )
            full_text_chunks.append(raw_text)
            current_offset = offset_end + 1  # newline spacing
            idx += 1

        # Also extract table text
        for t_idx, table in enumerate(doc.tables, start=1):
            table_rows_text: list[str] = []
            for row in table.rows:
                row_cells = [cell.text.strip() for cell in row.cells]
                row_str = " | ".join(c for c in row_cells if c)
                if row_str:
                    table_rows_text.append(row_str)

            if table_rows_text:
                table_block_text = f"[表格 {t_idx}]\n" + "\n".join(table_rows_text)
                offset_start = current_offset
                offset_end = offset_start + len(table_block_text)
                paragraphs.append(
                    ParagraphBlock(
                        id=f"p{idx}",
                        index=idx,
                        section=f"表格 {t_idx}",
                        text=table_block_text,
                        offset_start=offset_start,
                        offset_end=offset_end,
                    )
                )
                full_text_chunks.append(table_block_text)
                current_offset = offset_end + 1
                idx += 1

        full_text = "\n".join(full_text_chunks)
        return ParsedDocument(
            parsed_text=full_text,
            paragraphs=[asdict(p) for p in paragraphs],
            sha256=sha256,
            is_scanned=len(full_text.strip()) == 0,
        )

    def _parse_pdf(self, data: bytes, sha256: str) -> ParsedDocument:
        reader = PdfReader(io.BytesIO(data))
        paragraphs: list[ParagraphBlock] = []
        full_text_chunks: list[str] = []
        current_offset = 0
        idx = 1

        for page_num, page in enumerate(reader.pages, start=1):
            page_text = page.extract_text() or ""
            lines = [line.strip() for line in page_text.split("\n") if line.strip()]

            current_para_lines: list[str] = []
            for line in lines:
                current_para_lines.append(line)
                # Group lines ending with period, colon, semicolon or Chinese punctuation
                if line.endswith((".", "。", "；", ";", "：", ":", "！", "!")):
                    para_text = " ".join(current_para_lines)
                    offset_start = current_offset
                    offset_end = offset_start + len(para_text)
                    paragraphs.append(
                        ParagraphBlock(
                            id=f"p{idx}",
                            index=idx,
                            section=f"第 {page_num} 页",
                            text=para_text,
                            offset_start=offset_start,
                            offset_end=offset_end,
                        )
                    )
                    full_text_chunks.append(para_text)
                    current_offset = offset_end + 1
                    idx += 1
                    current_para_lines = []

            if current_para_lines:
                para_text = " ".join(current_para_lines)
                offset_start = current_offset
                offset_end = offset_start + len(para_text)
                paragraphs.append(
                    ParagraphBlock(
                        id=f"p{idx}",
                        index=idx,
                        section=f"第 {page_num} 页",
                        text=para_text,
                        offset_start=offset_start,
                        offset_end=offset_end,
                    )
                )
                full_text_chunks.append(para_text)
                current_offset = offset_end + 1
                idx += 1

        full_text = "\n".join(full_text_chunks)
        is_scanned = len(full_text.strip()) < 50

        return ParsedDocument(
            parsed_text=full_text,
            paragraphs=[asdict(p) for p in paragraphs],
            sha256=sha256,
            is_scanned=is_scanned,
        )

    def _parse_text(self, text: str, sha256: str) -> ParsedDocument:
        lines = [line.strip() for line in text.split("\n") if line.strip()]
        paragraphs: list[ParagraphBlock] = []
        full_text_chunks: list[str] = []
        current_offset = 0

        for idx, line in enumerate(lines, start=1):
            offset_start = current_offset
            offset_end = offset_start + len(line)
            paragraphs.append(
                ParagraphBlock(
                    id=f"p{idx}",
                    index=idx,
                    section="正文",
                    text=line,
                    offset_start=offset_start,
                    offset_end=offset_end,
                )
            )
            full_text_chunks.append(line)
            current_offset = offset_end + 1

        full_text = "\n".join(full_text_chunks)
        return ParsedDocument(
            parsed_text=full_text,
            paragraphs=[asdict(p) for p in paragraphs],
            sha256=sha256,
            is_scanned=len(full_text.strip()) == 0,
        )
