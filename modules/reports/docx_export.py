"""Deterministic Markdown → DOCX export for analysis reports.

The report pipeline stores reports as Markdown produced by
``modules.reports.generator.MarkdownReportGenerator``. This module converts that
Markdown into a Word document so agencies can deliver an editable file, as the
MVP spec requires ("PDF/DOCX 报告").

Scope is deliberately limited to the constructs the generator actually emits:
``#``/``##``/``###`` headings, pipe tables with ``:---`` separator rows,
``- `` bullets, ``> `` quotes, ``---`` rules, fenced ``` code blocks, and inline
``**bold**`` / `` `code` ``. Anything else falls back to plain paragraph text —
content is carried through verbatim; the converter adds no wording of its own
(no conclusion language can sneak in here, and the report's own
disclaimers/gates travel with the content).

python-docx is an existing project dependency; nothing new is introduced.

Byte-level determinism: python-docx stamps every zip entry with the wall-clock
time, so two exports of the same report would hash differently. When
provenance is supplied, the package timestamps and core properties are pinned
to the snapshot's ``sealed_at`` so the same sealed snapshot always yields the
same bytes — and therefore the same file SHA-256 (spec §4.1 "对最终文件计算
SHA-256"). The file hash itself cannot live inside the file; callers publish it
alongside (response header, audit record).
"""

from __future__ import annotations

import io
import re
import zipfile
from dataclasses import dataclass
from datetime import datetime, timezone

from docx import Document
from docx.shared import Pt

_HEADING_RE = re.compile(r"^(#{1,3})\s+(.*)$")
_BULLET_RE = re.compile(r"^-\s+(.*)$")
_QUOTE_RE = re.compile(r"^>\s?(.*)$")
_RULE_RE = re.compile(r"^---+\s*$")
_FENCE_RE = re.compile(r"^```(\w*)\s*$")
_TABLE_SEPARATOR_RE = re.compile(r"^\|?(\s*:?-+:?\s*\|)+\s*:?-+:?\s*\|?\s*$")
_INLINE_TOKEN_RE = re.compile(r"(\*\*.+?\*\*|`.+?`)")

# zip 格式可表示的最早时间；无溯源信息时作为固定时间戳
_ZIP_EPOCH = datetime(1980, 1, 1, tzinfo=timezone.utc)


@dataclass(frozen=True)
class ReportProvenance:
    """报告所属封存快照的溯源信息，写入页脚并用于固定文件时间戳。"""

    snapshot_number: int
    root_sha256: str
    sealed_at: datetime


def provenance_footer_text(provenance: ReportProvenance) -> str:
    """页脚文本：版本号 + 快照根校验值 + 封存时间（spec：版本号和校验值）。"""
    sealed = provenance.sealed_at.astimezone(timezone.utc).strftime(
        "%Y-%m-%dT%H:%M:%SZ"
    )
    return (
        f"证据快照版本 #{provenance.snapshot_number} · "
        f"快照根校验值 SHA-256: {provenance.root_sha256} · 封存时间 {sealed}"
    )


def _pin_zip_timestamps(data: bytes, moment: datetime) -> bytes:
    """以固定时间戳重写 zip 条目，使同一输入产出逐字节相同的文件。"""
    stamp = moment.astimezone(timezone.utc)
    if stamp < _ZIP_EPOCH:
        stamp = _ZIP_EPOCH
    date_time = stamp.timetuple()[:6]
    source = zipfile.ZipFile(io.BytesIO(data))
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as target:
        for info in source.infolist():
            pinned = zipfile.ZipInfo(info.filename, date_time=date_time)
            pinned.compress_type = zipfile.ZIP_DEFLATED
            pinned.external_attr = info.external_attr
            target.writestr(pinned, source.read(info.filename))
    return output.getvalue()


def _split_table_row(line: str) -> list[str]:
    """Split a pipe-table row into trimmed cell strings."""
    stripped = line.strip()
    if stripped.startswith("|"):
        stripped = stripped[1:]
    if stripped.endswith("|"):
        stripped = stripped[:-1]
    return [cell.strip() for cell in stripped.split("|")]


def _is_table_start(lines: list[str], index: int) -> bool:
    """A table needs a header row followed by a separator row."""
    if index + 1 >= len(lines):
        return False
    header, separator = lines[index].strip(), lines[index + 1].strip()
    return "|" in header and bool(_TABLE_SEPARATOR_RE.match(separator))


def _add_inline_runs(paragraph, text: str) -> None:
    """Add runs to a paragraph, honoring **bold** and `code` inline tokens."""
    for token in _INLINE_TOKEN_RE.split(text):
        if not token:
            continue
        if token.startswith("**") and token.endswith("**") and len(token) > 4:
            run = paragraph.add_run(token[2:-2])
            run.bold = True
        elif token.startswith("`") and token.endswith("`") and len(token) > 2:
            run = paragraph.add_run(token[1:-1])
            run.font.name = "Courier New"
        else:
            paragraph.add_run(token)


def markdown_to_docx_bytes(
    markdown: str,
    *,
    title: str = "专利证据分析报告",
    provenance: ReportProvenance | None = None,
) -> bytes:
    """把生成器产出的 Markdown 报告转换为 DOCX 字节流（纯函数、逐字节确定性）。"""
    document = Document()
    fixed_moment = provenance.sealed_at if provenance else _ZIP_EPOCH
    core = document.core_properties
    core.title = title
    core.author = "PatentEvidence"
    core.comments = ""
    core.created = fixed_moment.astimezone(timezone.utc).replace(tzinfo=None)
    core.modified = core.created
    core.last_modified_by = ""
    if provenance is not None:
        footer = document.sections[0].footer.paragraphs[0]
        footer.text = provenance_footer_text(provenance)
        for run in footer.runs:
            run.font.size = Pt(8)
    # 默认正文字号偏小，便于打印阅读
    normal = document.styles["Normal"]
    normal.font.size = Pt(10.5)

    document.add_heading(title, level=0)

    lines = markdown.splitlines()
    index = 0
    while index < len(lines):
        raw = lines[index]
        line = raw.strip()

        if not line:
            index += 1
            continue

        if _is_table_start(lines, index):
            header = _split_table_row(lines[index])
            index += 2  # 跳过表头与分隔行
            rows: list[list[str]] = []
            while index < len(lines) and "|" in lines[index].strip() and lines[index].strip():
                rows.append(_split_table_row(lines[index]))
                index += 1
            table = document.add_table(rows=1, cols=len(header))
            table.style = "Table Grid"
            for col, cell_text in enumerate(header):
                cell_paragraph = table.rows[0].cells[col].paragraphs[0]
                _add_inline_runs(cell_paragraph, cell_text)
                for run in cell_paragraph.runs:
                    run.bold = True
            for row in rows:
                cells = table.add_row().cells
                for col in range(len(header)):
                    cell_text = row[col] if col < len(row) else ""
                    _add_inline_runs(cells[col].paragraphs[0], cell_text)
            continue

        heading = _HEADING_RE.match(line)
        if heading:
            level = len(heading.group(1))
            document.add_heading(heading.group(2).strip(), level=min(level, 3))
            index += 1
            continue

        if _RULE_RE.match(line):
            # Word 没有水平线元素，用一个空段落隔开即可
            document.add_paragraph("")
            index += 1
            continue

        if _FENCE_RE.match(line):
            # 围栏代码块：逐行等宽渲染，直到闭合围栏（缺失则到文档尾）
            index += 1
            while index < len(lines) and not _FENCE_RE.match(lines[index].strip()):
                code_paragraph = document.add_paragraph()
                code_run = code_paragraph.add_run(lines[index])
                code_run.font.name = "Courier New"
                index += 1
            index += 1  # 跳过闭合围栏
            continue

        bullet = _BULLET_RE.match(line)
        if bullet:
            paragraph = document.add_paragraph(style="List Bullet")
            _add_inline_runs(paragraph, bullet.group(1))
            index += 1
            continue

        quote = _QUOTE_RE.match(line)
        if quote:
            paragraph = document.add_paragraph(style="Intense Quote")
            _add_inline_runs(paragraph, quote.group(1))
            index += 1
            continue

        paragraph = document.add_paragraph()
        _add_inline_runs(paragraph, line)
        index += 1

    buffer = io.BytesIO()
    document.save(buffer)
    return _pin_zip_timestamps(buffer.getvalue(), fixed_moment)
